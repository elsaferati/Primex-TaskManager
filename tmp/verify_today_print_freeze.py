"""Read-only hosted/database checks; no snapshot writes and no email delivery."""
from __future__ import annotations

import asyncio
import base64
import hashlib
import json
import sys
from datetime import datetime
from io import BytesIO
from pathlib import Path

import httpx
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import text

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
from app.config import settings
from app.db import SessionLocal, engine
from app.services.primeflow_report import PrimeFlowClient, report_timezone
from app.services.today_print_report_freeze import _source_payload, pack_report
from app.services.tomorrow_print_report import build_today_print_report, next_working_day


async def hosted() -> None:
    base = "https://api-flow.primexeu.com"
    async with httpx.AsyncClient(base_url=base, timeout=30) as client:
        response = await client.get("/health")
        print("hosted_health", response.status_code, response.json())
        login = PrimeFlowClient(base, settings.PRIMEFLOW_EMAIL or settings.ADMIN_EMAIL,
                               settings.PRIMEFLOW_PASSWORD or settings.ADMIN_PASSWORD,
                               settings.PRIMEFLOW_ACCESS_TOKEN)
        token = await login._token(client)
        for path in ("/api/today-print-report/freeze-status", "/api/today-print-report/frozen-history"):
            response = await client.get(path, headers={"Authorization": f"Bearer {token}"})
            body = response.json()
            print("hosted_endpoint", path, response.status_code, json.dumps(body, ensure_ascii=True)[:1500])


async def database() -> None:
    today = datetime.now(report_timezone()).date()
    async with SessionLocal() as db:
        await db.execute(text("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ READ ONLY"))
        rows = (await db.execute(text("""
            SELECT report_kind, report_date,
                   report_payload->>'frozen_at' AS frozen_at,
                   report_payload->>'source_captured_at' AS captured_at,
                   report_payload ? 'freeze_source' AS pending_render,
                   jsonb_array_length(COALESCE(report_payload->'frozen_attachments','[]'::jsonb)) AS artifacts
            FROM one_h_print_report_snapshots
            WHERE report_kind IN ('TODAY', 'TODAY_FROZEN')
            ORDER BY report_date DESC LIMIT 8
        """))).mappings().all()
        print("database_snapshots", json.dumps([dict(row) for row in rows], default=str))
        payload = await _source_payload(db, today)
        next_day = next_working_day(today)
        next_payload = await _source_payload(db, next_day) if next_day.isoformat() > payload["week_end"] else payload
        print("source_payload", "ok", {bucket: len(values) for bucket, values in payload["items"].items()})
        await db.rollback()
    # Generate a verification copy; it is never saved as the actual frozen report.
    report = await build_today_print_report(today, include_attachment=True, payload=payload,
                                           next_day_payload=next_payload)
    packed = pack_report(report)
    artifact_checks = []
    for original, stored in zip(report["attachments"], packed["frozen_attachments"], strict=True):
        filename, content, mime_type = original
        restored = base64.b64decode(stored["content"])
        assert hashlib.sha256(content).digest() == hashlib.sha256(restored).digest()
        if mime_type == "image/png":
            with Image.open(BytesIO(restored)) as im:
                im.verify()
            with Image.open(BytesIO(restored)) as im:
                artifact_checks.append({"file": filename, "dimensions": im.size, "roundtrip": "ok"})
        else:
            book = load_workbook(BytesIO(restored))
            artifact_checks.append({"file": filename, "rows": book.active.max_row, "roundtrip": "ok"})
    print("report_exports", json.dumps(artifact_checks))
    assert "<table" in report["html"] and "background" in report["html"]
    print("report_html", "tables_and_colors_ok")


async def main() -> None:
    try:
        if "--hosted" in sys.argv:
            await hosted()
        if "--database" in sys.argv:
            await database()
    finally:
        await engine.dispose()


if __name__ == "__main__":
    asyncio.run(main())
