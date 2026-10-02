from __future__ import annotations

import unittest
import uuid
from datetime import date, datetime, timezone
from email import policy
from email.parser import BytesParser
from io import StringIO
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

from alembic.migration import MigrationContext
from alembic.operations import Operations
from fastapi import FastAPI
from fastapi.testclient import TestClient
from pydantic import ValidationError

from app.api.routers import m3_reporting_points as api
from app.models.enums import UserRole
from app.models.m3_reporting_points import M3ReportingPointsReport
from app.services import m3_reporting_points as service
from app.services import m3_reporting_points_scheduler as scheduler
from app.services.primeflow_report import GmailService, REPORT_SENDER_EMAIL
from tests.test_migration_graph import migration_scripts

DAY = date(2026, 10, 5)


def task(**changes):
    values = dict(id=uuid.uuid4(), title="Detyre", status="TODO", progress_percentage=0,
                  created_at=datetime(2026, 10, 5, 6, tzinfo=timezone.utc),
                  start_date=datetime(2026, 10, 5, 7, tzinfo=timezone.utc),
                  due_date=datetime(2026, 10, 5, 14, tzinfo=timezone.utc),
                  is_deadline_important=False, completed_at=None)
    return SimpleNamespace(**{**values, **changes})


def event(item, field="due_date", before="2026-10-05", after="2026-10-06", **changes):
    values = dict(id=uuid.uuid4(), entity_id=item.id, action=f"task.{field}_changed",
                  before={"value": before}, after={"value": after},
                  created_at=datetime(2026, 10, 5, 8, tzinfo=timezone.utc))
    return SimpleNamespace(**{**values, **changes})


class SelectionTests(unittest.TestCase):
    def test_todo_requires_today_within_start_due_interval_even_with_prior_progress(self):
        rows = [task(title="Me progres", progress_percentage=65),
                task(title="Intervali", start_date=date(2026, 10, 2), due_date=date(2026, 10, 8)),
                task(title="E vjeter", due_date=datetime(2026, 9, 1, tzinfo=timezone.utc)),
                task(title="Per te ardhmen", start_date=datetime(2026, 11, 1, tzinfo=timezone.utc)),
                task(title="Pa Start", start_date=None), task(title="Pa Due", due_date=None),
                task(title="DONE", status="DONE"), task(title="IN_PROGRESS", status="IN_PROGRESS")]
        sections = service.select_task_sections(rows, [], DAY)
        self.assertEqual([row["title"] for row in sections["untouched"]], ["Me progres", "Intervali"])

    def test_postponed_uses_original_week_interval_excluding_future_and_past_weeks(self):
        rows = [task(title="Kjo jave", start_date=date(2026, 10, 12), due_date=date(2026, 10, 14)),
                task(title="Mbulon javen", start_date=date(2026, 9, 28), due_date=date(2026, 10, 14)),
                task(title="Future", start_date=date(2026, 10, 12), due_date=date(2026, 10, 15)),
                task(title="E kaluar", start_date=date(2026, 9, 21), due_date=date(2026, 10, 7)),
                task(title="Pa shtyrje")]
        events = [event(rows[0], before="2026-10-09", after="2026-10-14"),
                  event(rows[0], "start_date", before="2026-10-05", after="2026-10-12"),
                  event(rows[1], before="2026-10-12", after="2026-10-14"),
                  event(rows[2], before="2026-10-14", after="2026-10-15"),
                  event(rows[3], before="2026-09-25", after="2026-10-07")]
        result = service.select_task_sections(rows, events, DAY)["postponed"]
        self.assertEqual([row["title"] for row in result], ["Kjo jave", "Mbulon javen"])

    def test_task_title_uses_existing_report_cleanup_without_description(self):
        row = task(title="[[added]]DT: Titulli i detyres due 16:00[[/added]]\n1. Pershkrimi i plote\n2. Hapi tjeter")
        selected = service.select_task_sections([row], [], DAY)
        self.assertEqual(selected["untouched"][0]["title"], "DT: Titulli i detyres")

    def test_start_postponement_within_report_week_is_ok_for_an_ongoing_task(self):
        row = task(start_date=date(2026, 10, 6), due_date=date(2026, 10, 8))
        changed = event(row, "start_date", before="2026-09-30", after="2026-10-06")
        result = service.select_task_sections([row], [changed], DAY)["postponed"]
        self.assertEqual(result[0]["risk"], "OK")

    def test_risk_distinguishes_within_week_friday_and_next_week(self):
        for target, expected in ((date(2026, 10, 6), "OK"), (date(2026, 10, 9), "RREZIK"),
                                 (date(2026, 10, 12), "RREZIK"), (date(2026, 10, 20), "RREZIK")):
            with self.subTest(target=target):
                row = task(due_date=target)
                result = service.select_task_sections([row], [event(row, after=target.isoformat())], DAY)
                self.assertEqual(result["postponed"][0]["risk"], expected)
                self.assertEqual(result["postponed"][0]["risk_color"], "#dcfce7" if expected == "OK" else "#fee2e2")

    def test_postponements_identify_due_only_both_dates_and_start_only(self):
        due = task(due_date=date(2026, 10, 6))
        both = task(start_date=date(2026, 10, 6), due_date=date(2026, 10, 7))
        start = task(start_date=date(2026, 10, 6), due_date=date(2026, 10, 8))
        events = [event(due), event(both, "start_date"), event(both, after="2026-10-07"), event(start, "start_date")]
        rows = service.select_task_sections([due, both, start], events, DAY)["postponed"]
        self.assertEqual([row["postponement_kind"] for row in rows], ["due", "start_due", "start"])
        groups = service.task_table_groups(rows, "postponed")
        self.assertEqual([len(group[2]) for group in groups], [1, 1, 1])
        self.assertEqual(len({row["task_id"] for _, _, rows in groups for row in rows}), 3)

    def test_am_pm_and_task_type_use_the_existing_m3_rules(self):
        rows = [task(system_task_slot_id=uuid.uuid4(), finish_period="AM"),
                task(is_1h_report=True, finish_period="PM"), task(is_bllok=True, finish_period="AM/PM"),
                task(project_id=uuid.uuid4()), task(is_personal=True)]
        result = service.select_task_sections(rows, [], DAY)["untouched"]
        self.assertEqual([row["task_type"] for row in result], ["SYS", "1H", "BLL", "PRJK", "P"])
        self.assertEqual([row["am_pm"] for row in result], ["AM", "PM", "AM/PM", "-", "-"])

    def test_task_symbols_and_eight_am_metadata_follow_existing_task_rules(self):
        rows = [task(one_h_marker="MONITOR", one_h_marker_by_ga=True, one_h_marker_date=date(2024, 1, 1), title="Task monitor"),
                task(one_h_marker="EXCLAMATION", title="EM: Task"),
                task(system_task_slot_id=uuid.uuid4(), title="EM: SYS"),
                task(due_date=datetime(2026, 10, 5, 6, tzinfo=timezone.utc), title="Task due at eight")]
        result = service.select_task_sections(rows, [], DAY)["untouched"]
        self.assertEqual([row["marker"] for row in result], ["(👁)", "!", "", ""])
        self.assertEqual([row["eight_am"] for row in result], [False, True, False, True])
        self.assertTrue(result[2]["is_system_task"])

    def test_same_day_requires_all_three_dates_and_excludes_done(self):
        rows = [task(status=status, progress_percentage=percent) for status, percent in (
            ("TODO", 0), ("IN_PROGRESS", 35), ("DONE", 100), ("WAITING_CLIENT", 15),
        )]
        rows += [task(created_at=date(2026, 10, 4)), task(start_date=date(2026, 10, 4)), task(due_date=date(2026, 10, 6))]
        result = service.select_task_sections(rows, [], DAY)["same_day"]
        self.assertEqual([(row["status"], row["progress"]) for row in result],
                         [("TODO", 0), ("IN_PROGRESS", 35), ("WAITING_CLIENT", 15)])

    def test_ga_uses_deadline_important_and_not_merely_a_due_date(self):
        plain = task(created_at=date(2026, 9, 1), due_date=date(2026, 10, 6))
        deadline = task(created_at=date(2026, 9, 1), due_date=date(2026, 10, 6), is_deadline_important=True)
        same_day = task(start_date=date(2026, 10, 6), due_date=date(2026, 10, 6))
        both = task(due_date=date(2026, 10, 6), is_deadline_important=True)
        not_moved = task(is_deadline_important=True)
        events = [event(plain), event(deadline), event(same_day),
                  event(same_day, "start_date"), event(both)]
        result = service.select_task_sections([plain, deadline, same_day, both, not_moved], events, DAY)["ga_postponed"]
        self.assertEqual([row["task_id"] for row in result], [str(deadline.id), str(same_day.id), str(both.id)])
        self.assertEqual(result[-1]["category"], "Deadline Important / SOT/SOT")

    def test_creation_uses_the_local_day_and_initial_date_assignment_is_not_a_move(self):
        row = task(created_at=datetime(2026, 10, 4, 22, 30, tzinfo=timezone.utc), due_date=date(2026, 10, 6))
        initial = event(row, before=None, after="2026-10-05", created_at=datetime(2026, 10, 5, 7, tzinfo=timezone.utc))
        moved = event(row)
        result = service.select_task_sections([row], [initial, moved], DAY)
        self.assertEqual(len(result["ga_postponed"]), 1)
        self.assertEqual(result["ga_postponed"][0]["created_date"], "2026-10-05")

    def test_reverted_or_other_day_changes_do_not_count_as_postponed(self):
        row = task(is_deadline_important=True)
        changes = [event(row), event(row, before="2026-10-06", after="2026-10-05")]
        result = service.select_task_sections([row], changes, DAY)
        self.assertEqual(result["postponed"], [])
        self.assertEqual(result["ga_postponed"], [])
        row.due_date = date(2026, 10, 6)
        yesterday = event(row, created_at=datetime(2026, 10, 4, 8, tzinfo=timezone.utc))
        self.assertEqual(service.select_task_sections([row], [yesterday], DAY)["ga_postponed"], [])

    def test_ga_checks_dates_immediately_before_postponement_not_only_day_start(self):
        row = task(due_date=date(2026, 10, 6))
        normalize_start = event(row, "start_date", before="2026-10-04", after="2026-10-05",
                                created_at=datetime(2026, 10, 5, 7, tzinfo=timezone.utc))
        result = service.select_task_sections([row], [normalize_start, event(row)], DAY)["ga_postponed"]
        self.assertEqual(len(result), 1)
        self.assertEqual(result[0]["category"], "SOT/SOT")
        self.assertEqual(result[0]["old_start_date"], "2026-10-05")

    def test_completed_late_date_normalization_is_not_a_postponement(self):
        row = task(status="DONE", is_deadline_important=True, completed_at=datetime(2026, 10, 5, 12, tzinfo=timezone.utc))
        result = service.select_task_sections([row], [event(row, before="2026-10-02", after="2026-10-05")], DAY)
        self.assertEqual(result["ga_postponed"], [])


class RenderingTests(unittest.TestCase):
    def test_system_split_deadline_color_eight_am_outline_markers_and_department_summary(self):
        report = self.report()
        report["data"] = {"untouched": [
            {"task_id": "regular", "title": "Task priority", "status": "TODO", "task_type": "1H", "deadline_important": True, "marker": "!!!"},
            {"task_id": "system", "title": "08:00 Task SYS", "status": "TODO", "task_type": "SYS", "marker": "👁"},
        ]}
        report["realization"]["departments"] = [{"department_id": "dev", "code": "DEV", "percent": 75, "comment": "Mbi 50%"}]
        html = service.render_html(report)
        self.assertIn("DET FT DHE PRJK PA PROGRES:", html)
        self.assertIn("DETYRAT E SISTEMIT PA PROGRES:", html)
        self.assertIn("bgcolor='#dc2626'", html)
        self.assertIn("border-top:3px solid #dc2626", html)
        self.assertIn("data-task-symbol='true'", html)
        self.assertIn("👁</strong>", html)
        self.assertIn("!!!</strong>", html)
        self.assertIn("75%", html)
        self.assertIn("DEV: 75% - Mbi 50%", service.render_plain_text(report))
        self.assertIn("👁 08:00 Task SYS", service.render_plain_text(report))

    def report(self):
        rows = []
        for target in (date(2026, 10, 6), date(2026, 10, 9)):
            row = task(title="<script>alert('x')</script>", due_date=target)
            moved = service.select_task_sections([row], [event(row, after=target.isoformat())], DAY)["postponed"][0]
            moved.update(comment="Klienti kerkoi shtyrjen\nDuhet informacion", reason="Ne pritje te klientit", assignees="EF", department="DEV", project="P")
            rows.append(moved)
        return {"report_date": DAY.isoformat(), "manual_answers": {"underload": "<b>Pergjigje</b>"},
                "data": {"postponed": rows}, "realization": {"percent": 62.0, "comment": service.realization_comment(62)},
                "realization_captured_at": "2026-10-05T16:15:00+02:00"}

    def test_email_retains_daily_reasons_comments_colors_manual_answers_and_both_blocks(self):
        report = self.report()
        result = service.render_html(report)
        self.assertIn(service.M3_TITLE, result)
        self.assertIn(service.GA_TITLE, result)
        self.assertIn("background:#fee2e2", result)
        self.assertIn("background:#dcfce7", result)
        self.assertIn("Ne pritje te klientit", result)
        self.assertIn("Klienti kerkoi shtyrjen<br>Duhet informacion", result)
        self.assertIn("62%", result)
        self.assertNotIn("<script>", result)
        self.assertIn("&lt;b&gt;Pergjigje&lt;/b&gt;", result)
        self.assertIn("Ne pritje te klientit", service.render_plain_text(report))
        self.assertIn("Klienti kerkoi shtyrjen", service.render_plain_text(report))

    def test_threshold_and_missing_data_are_not_invented_zero(self):
        self.assertIn("mbi 50%", service.realization_comment(62))
        self.assertIn("nen 50%", service.realization_comment(42))
        self.assertEqual(service.realization_comment(50), "Jemi ne 50%.")
        self.assertIn("Mungojne", service.realization_comment(None))

    def test_tables_have_uppercase_headers_status_colors_and_no_progress_columns(self):
        report = self.report()
        report["data"]["same_day"] = [{**report["data"]["postponed"][0], "status": "IN_PROGRESS", "progress": 35,
                                       "title": "DT: Vetem titulli\n1. Pershkrimi", "completed_today": 999}]
        report["data"]["postponed"].append({**report["data"]["postponed"][0], "status": "DONE"})
        result = service.render_html(report)
        self.assertIn(">TITULLI</th>", result)
        self.assertIn(">AM/PM</th>", result)
        self.assertIn(">LLOJI</th>", result)
        self.assertNotIn(">STATUSI</th>", result)
        self.assertIn(">KOMENTI</th>", result)
        self.assertIn("background:#FFC4ED", result)
        self.assertIn("background:#C4FDC4", result)
        self.assertIn("DT: Vetem titulli", result)
        self.assertNotIn("Pershkrimi", result)
        self.assertNotIn("Progresi %", result)
        self.assertNotIn("Progresi sot", result)
        self.assertNotIn("999", result)
        plain = service.render_plain_text(report)
        self.assertNotIn("progress:", plain)
        self.assertNotIn("completed_today:", plain)
        self.assertNotIn("Pershkrimi", plain)

    def test_split_tables_render_stacked_start_due_in_from_to_cells(self):
        report = self.report()
        report["data"]["postponed"][0].update(postponement_kind="start_due", old_start_date="2026-10-05", start_date="2026-10-06")
        result = service.render_html(report)
        self.assertIn("SHTYER START DHE DUE DATE:", result)
        self.assertIn("SHTYER DUE DATE:", result)
        self.assertIn(">NGA</th>", result)
        self.assertIn(">NE</th>", result)
        self.assertIn("START: 05.10.2026</td>", result)
        self.assertIn("DUE: 05.10.2026</td>", result)
        self.assertIn("border-bottom:3px solid #000", result)
        self.assertIn("role='presentation'", result)
        self.assertIn("table-layout:fixed", result)
        self.assertIn("<col style='width:28px'>", result)
        self.assertIn("<col style='width:42px'>", result)
        self.assertIn("<col><col style='width:124px'>", result)

    def test_untouched_and_same_day_omit_dates_and_same_day_excludes_done(self):
        rows = [{"task_id": "done", "title": "Done title", "status": "DONE", "created_date": "2026-10-05", "start_date": "2026-10-05", "due_date": "2026-10-05"},
                {"task_id": "active", "title": "Active title", "status": "IN_PROGRESS"},
                {"task_id": "todo", "title": "Todo title", "status": "TODO"}]
        report = self.report()
        report["data"] = {"same_day": rows}
        result = service.render_html(report)
        self.assertNotIn(">CREATION / START / DUE</th>", result)
        self.assertNotIn("CREATION:", result)
        self.assertNotIn("START:", result)
        self.assertNotIn("DUE:", result)
        self.assertLess(result.index("Active title"), result.index("Todo title"))
        self.assertNotIn("Done title", result)
        self.assertNotIn("Done title", service.render_plain_text(report))
        self.assertNotIn("CREATION / START / DUE:", service.render_plain_text(report))
        row = M3ReportingPointsReport(id=uuid.uuid4(), report_date=DAY, data=report["data"],
                                      manual_answers={}, realization={"percent": 62}, status="SENT")
        payload = service.report_payload(row)
        self.assertEqual([item["task_id"] for item in payload["data"]["same_day"]], ["active", "todo"])
        self.assertEqual(len(row.data["same_day"]), 3)  # Original saved copy is preserved.
        self.assertEqual(payload["realization"], {"percent": 62})
        report["data"] = {"untouched": [rows[-1]]}
        result = service.render_html(report)
        self.assertNotIn(">START</th>", result)
        self.assertNotIn(">DUE</th>", result)

    def test_manual_payload_rejects_automatic_fields_and_keeps_newlines(self):
        self.assertEqual(api.ManualAnswersPayload(manual_answers={"reorganization": "Rreshti 1\nRreshti 2"}).manual_answers["reorganization"], "Rreshti 1\nRreshti 2")
        with self.assertRaises(ValidationError):
            api.ManualAnswersPayload(manual_answers={"percent": "100"})
        with self.assertRaises(ValidationError):
            api.ManualAnswersPayload(manual_answers={"underload": "x" * 10001})

    def test_migration_emits_unique_per_date_table_and_no_auto_send_settings(self):
        output = StringIO()
        migration = migration_scripts().get_revision("0142_m3_reporting_points").module
        operations = Operations(MigrationContext.configure(dialect_name="postgresql", opts={"as_sql": True, "output_buffer": output}))
        with patch.object(migration, "op", operations):
            migration.upgrade()
        sql = output.getvalue()
        self.assertIn("CREATE TABLE m3_reporting_points_reports", sql)
        self.assertIn("UNIQUE (report_date)", sql)
        self.assertIn("realization_captured_at TIMESTAMP WITH TIME ZONE", sql)


class CaptureAndWorkflowTests(unittest.IsolatedAsyncioTestCase):
    async def test_build_tables_reuses_weekly_planner_department_user_order_and_short_labels(self):
        departments = [SimpleNamespace(id=uuid.uuid4(), code=code) for code in ("DEV", "GD", "PCM")]
        # Alphabetical order differs from the saved Weekly Planner order.
        users = [SimpleNamespace(id=uuid.uuid4(), full_name=name, department_id=departments[dep].id,
                                 weekly_planner_sort_order=order) for name, dep, order in
                 (("Zoe Dev", 0, 1), ("Adam Dev", 0, 2), ("Genti Design", 1, 1), ("Elsa Product", 2, 1))]
        tasks = [task(title=user.full_name, assigned_to=user.id, department_id=departments[2].id,
                      project_id=None, fast_task_order=0, finish_period="PM") for user in users]
        def result(values):
            item = MagicMock()
            item.scalars.return_value.all.return_value = values
            item.all.return_value = values
            return item
        for mode in ("active", "postponed", "done"):
            events = []
            for item in tasks:
                item.due_date = DAY
            tasks[0].status = "DONE" if mode == "done" else "TODO"
            if mode == "postponed":
                for item in tasks:
                    item.due_date = date(2026, 10, 6)
                    events.append(event(item, after="2026-10-06"))
                    events[-1].after["reason"] = "Arsyeja e ndryshimit te dates"
            daily = [SimpleNamespace(task_id=tasks[0].id, user_id=users[0].id,
                                     reason_code="WAITING_CLIENT", comment="Pres pergjigjen e klientit")]
            db = AsyncMock()
            db.execute.side_effect = [result(list(reversed(tasks))), result(events), result(users), result(departments),
                                      result([]), result([]), result(users), result(daily), result([])]
            data = await service.build_task_data(db, DAY)
            for key in (("untouched", "postponed", "ga_postponed") if mode == "postponed" else ("untouched", "same_day")):
                with self.subTest(key=key):
                    expected = ["ZD", "AD", "GD", "EP"] if mode != "done" else ["AD", "GD", "EP"]
                    self.assertEqual([row["assignees"] for row in data[key]], expected)
                    for row in data[key]:
                        if row["assignees"] == "ZD":
                            self.assertEqual(row["department"], "DEV")
                            self.assertEqual(row["reason"], service.REASON_LABELS["WAITING_CLIENT"])
                            self.assertEqual(row["comment"], "Pres pergjigjen e klientit")
                        else:
                            self.assertEqual(row["reason"], "-")
                            self.assertEqual(row["comment"], "-")
                    if mode != "done":
                        self.assertEqual([row["department"] for row in data[key]], ["DEV", "DEV", "GD", "PCM"])
            query = db.execute.await_args_list[7].args[0]
            self.assertIn("task_daily_rlz_states.day_date", str(query))
            self.assertIn(DAY, query.compile().params.values())

    async def test_capture_runs_only_at_1615_on_weekdays_and_never_late(self):
        with patch.object(scheduler, "SessionLocal") as session:
            for stamp in ("2026-10-05T16:14:59+02:00", "2026-10-05T16:16:00+02:00", "2026-10-03T16:15:00+02:00"):
                self.assertFalse(await scheduler.run_m3_reporting_points_scheduler_once(datetime.fromisoformat(stamp)))
        session.assert_not_called()
        self.assertTrue(scheduler.is_capture_minute(datetime.fromisoformat("2026-10-05T14:15:30+00:00")))

    async def test_capture_is_immutable_preserves_manual_answers_and_sends_no_email(self):
        row = SimpleNamespace(realization_captured_at=None, realization=None, generated_at=datetime.now(timezone.utc),
                              status="DRAFT", manual_answers={"underload": "Pergjigje e ruajtur"})
        db = AsyncMock()
        context = MagicMock()
        context.__aenter__ = AsyncMock(return_value=db)
        context.__aexit__ = AsyncMock(return_value=False)
        with (patch.object(scheduler, "SessionLocal", return_value=context),
              patch.object(scheduler, "locked_report", new=AsyncMock(return_value=row)),
              patch.object(scheduler, "build_realization_capture", new=AsyncMock(return_value={"percent": 62})) as capture,
              patch.object(service, "GmailService") as gmail):
            self.assertTrue(await scheduler.run_m3_reporting_points_scheduler_once(datetime.fromisoformat("2026-10-05T16:15:10+02:00")))
            self.assertFalse(await scheduler.run_m3_reporting_points_scheduler_once(datetime.fromisoformat("2026-10-05T16:15:40+02:00")))
        capture.assert_awaited_once()
        gmail.assert_not_called()
        self.assertEqual(row.manual_answers, {"underload": "Pergjigje e ruajtur"})
        self.assertEqual(row.realization, {"percent": 62})

    async def test_staff_percentage_is_weighted_by_the_existing_realization_metric(self):
        department = uuid.uuid4()
        ids = [uuid.uuid4(), uuid.uuid4()]
        users = [SimpleNamespace(id=user_id, department_id=department) for user_id in ids]
        result = MagicMock()
        result.scalars.return_value.all.return_value = users
        department_result = MagicMock()
        department_result.scalars.return_value.all.return_value = [SimpleNamespace(id=department, code="DEV", name="Development")]
        db = AsyncMock()
        db.execute.side_effect = [result, department_result]
        # One person realizes 1/1, the other 0/3: staff result must be 25%, not 50%.
        people = [{"user_id": str(ids[0]), "user_name": "A", "tasks": [{"in_original_plan": True, "classification": "REALIZED_AS_PLANNED"}], "metrics": {"raw_plan_realization": 100}},
                  {"user_id": str(ids[1]), "user_name": "B", "tasks": [{"in_original_plan": True, "classification": "NO_PROGRESS"}] * 3, "metrics": {"raw_plan_realization": 0}}]
        with patch.object(service, "build_live_daily_realization", new=AsyncMock(return_value={"baseline_available": True, "people": people})):
            capture = await service.build_realization_capture(db, DAY)
        self.assertEqual(capture["percent"], 25)
        self.assertEqual(capture["departments"][0]["code"], "DEV")
        self.assertEqual(capture["departments"][0]["percent"], 25)
        db.execute.side_effect = [result, department_result]
        with patch.object(service, "build_live_daily_realization", new=AsyncMock(return_value={"baseline_available": False, "people": people})):
            capture = await service.build_realization_capture(db, DAY)
        self.assertIsNone(capture["percent"])
        self.assertIsNone(capture["departments"][0]["percent"])

    async def test_departments_have_separate_weighted_percentages_and_saved_order(self):
        dev, gd, pcm = [SimpleNamespace(id=uuid.uuid4(), code=code, name=code) for code in ("DEV", "GD", "PCM")]
        users = [SimpleNamespace(id=uuid.uuid4(), department_id=item.id) for item in (dev, gd, pcm)]
        def result(values):
            item = MagicMock()
            item.scalars.return_value.all.return_value = values
            return item
        db = AsyncMock()
        db.execute.side_effect = [result(users), result([pcm, gd, dev])]
        live = {dev.id: {"baseline_available": True, "people": [{"user_id": str(users[0].id), "tasks": [{"in_original_plan": True, "classification": "REALIZED_AS_PLANNED"}], "metrics": {}}]},
                gd.id: {"baseline_available": True, "people": [{"user_id": str(users[1].id), "tasks": [{"in_original_plan": True, "classification": "NO_PROGRESS"}] * 3, "metrics": {}}]},
                pcm.id: {"baseline_available": False, "people": []}}
        async def build(_, department_id, day):
            return live[department_id]
        with patch.object(service, "build_live_daily_realization", new=build):
            capture = await service.build_realization_capture(db, DAY)
        self.assertEqual([(item["code"], item["percent"]) for item in capture["departments"]], [("DEV", 100), ("GD", 0), ("PCM", None)])
        self.assertIsNone(capture["percent"])
        self.assertEqual([item["employees"] for item in capture["departments"]], [1, 1, 1])

    async def test_regeneration_preserves_answers_and_realization(self):
        row = SimpleNamespace(report_date=datetime.now(service.report_timezone()).date(),
                              manual_answers={"ga_reorganization": "Jo"}, realization={"percent": 62},
                              realization_captured_at="16:15", data={}, generated_at=None)
        with patch.object(service, "build_task_data", new=AsyncMock(return_value={"untouched": [{"title": "TODO"}]})):
            await service.refresh_report(AsyncMock(), row)
        self.assertEqual(row.manual_answers, {"ga_reorganization": "Jo"})
        self.assertEqual(row.realization, {"percent": 62})
        self.assertEqual(row.realization_captured_at, "16:15")

    async def test_historical_regeneration_never_reads_current_task_status(self):
        row = SimpleNamespace(report_date=date(2026, 1, 1))
        with patch.object(service, "build_task_data", new=AsyncMock()) as builder:
            with self.assertRaises(ValueError):
                await service.refresh_report(AsyncMock(), row)
        builder.assert_not_awaited()

    async def test_sent_report_is_idempotent_and_blank_recipients_do_not_send(self):
        row = SimpleNamespace(status="SENT")
        with patch.object(service, "GmailService") as gmail:
            await service.send_report(AsyncMock(), row, {"to": ["recipient@example.com"]})
            row.status = "DRAFT"
            with self.assertRaises(ValueError):
                await service.send_report(AsyncMock(), row, {"to": []})
        gmail.assert_not_called()

    async def test_manual_send_attaches_the_entire_report_without_real_email(self):
        row = M3ReportingPointsReport(id=uuid.uuid4(), report_date=DAY, manual_answers={"underload": "Koment"},
                                      data={"untouched": []}, realization={"percent": 62, "comment": "Mbi 50%"},
                                      realization_captured_at=datetime.fromisoformat("2026-10-05T16:15:00+02:00"),
                                      status="DRAFT")
        gmail = SimpleNamespace(find_exact=AsyncMock(return_value=None), send_verified=AsyncMock(return_value={"id": "message"}))
        db = AsyncMock()
        recipients = {"to": ["current-m3@example.com"], "cc": ["cc@example.com"], "bcc": []}
        with patch.object(service, "GmailService", return_value=gmail):
            await service.send_report(db, row, recipients)
        call = gmail.send_verified.await_args
        self.assertEqual(call.args[1], recipients)
        self.assertEqual(call.kwargs["attachments"][0][1].decode(), call.args[3])
        self.assertEqual(call.kwargs["attachments"][0][2], "text/html")
        self.assertEqual([a[2] for a in call.kwargs["attachments"]], ["text/html", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "image/png"])
        self.assertEqual(row.status, "SENT")
        db.commit.assert_awaited_once()

    async def test_smtp_email_preserves_full_report_colors_thick_dividers_and_html_attachment(self):
        statuses = ["TODO", "IN_PROGRESS", "WAITING_CLIENT", "WAITING_CONFIRMATION", "DONE"]
        rows = [{"task_id": str(uuid.uuid4()), "title": f"Detyrë {status}", "status": status,
                 "assignees": "EF", "department": "DEV", "am_pm": "PM", "task_type": "1H",
                 "old_start_date": "2026-10-05", "old_due_date": "2026-10-05",
                 "start_date": "2026-10-06", "due_date": "2026-10-09",
                 "postponement_kind": "start_due", "risk": "OK" if index == 0 else "RREZIK",
                 "reason": "Në pritje të klientit", "comment": "Klienti kërkoi shtyrjen\nRreshti tjetër"}
                for index, status in enumerate(statuses)]
        # The complete attachment must survive inline email display limits.
        many_todo = [{**rows[0], "task_id": str(uuid.uuid4()), "title": f"Detyrë e paprekur {index}"} for index in range(120)]
        row = M3ReportingPointsReport(id=uuid.uuid4(), report_date=DAY, manual_answers={"underload": "Përgjigje me ë dhe ç"},
                                      data={"postponed": rows, "untouched": many_todo},
                                      realization={"percent": 62, "comment": "Jemi mbi 50%"},
                                      realization_captured_at=datetime.fromisoformat("2026-10-05T16:15:00+02:00"), status="DRAFT")
        recipients = {"to": ["current-m3@example.com"], "cc": ["cc@example.com"], "bcc": ["bcc@example.com"]}
        expected_html = service.render_html(service.report_payload(row))
        self.assertGreater(len(expected_html.encode("utf-8")), 102400)
        smtp = MagicMock()
        gmail = GmailService(sender=REPORT_SENDER_EMAIL, password="test-only-password")
        with (patch.object(service, "GmailService", return_value=gmail),
              patch("app.services.primeflow_report.smtplib.SMTP") as transport):
            transport.return_value.__enter__.return_value = smtp
            await service.send_report(AsyncMock(), row, recipients)
        smtp.send_message.assert_called_once()
        sent = smtp.send_message.call_args.args[0]
        message = BytesParser(policy=policy.default).parsebytes(sent.as_bytes())
        inline_html = message.get_body(preferencelist=("html",)).get_content()
        attachments = list(message.iter_attachments())
        self.assertEqual([a.get_content_type() for a in attachments], ["text/html", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "image/png"])
        from docx import Document
        from PIL import Image
        from io import BytesIO
        word = Document(BytesIO(attachments[1].get_payload(decode=True)))
        word_text = "\n".join(cell.text for table in word.tables for row in table.rows for cell in row.cells)
        self.assertIn("Detyrë e paprekur 119", word_text)
        self.assertIn("Përgjigje me ë dhe ç", "\n".join(p.text for p in word.paragraphs))
        png = Image.open(BytesIO(attachments[2].get_payload(decode=True)))
        png.verify()
        attachment = attachments[0]
        attached_html = attachment.get_payload(decode=True).decode("utf-8")
        self.assertEqual(inline_html.strip(), expected_html.strip())
        self.assertEqual(attached_html, expected_html)
        self.assertEqual(attachment.get_content_type(), "text/html")
        self.assertEqual(smtp.send_message.call_args.kwargs["to_addrs"], sum(recipients.values(), []))
        for color in (*service.STATUS_COLORS.values(), "#fee2e2", "#dcfce7"):
            self.assertIn(f"bgcolor='{color}'", inline_html)
        self.assertIn("border-bottom:3px solid #000", inline_html)
        self.assertIn("border:1px solid #000", inline_html)
        self.assertIn("background-color:#dcfce7;color:#14532d;padding:2px 4px'>Brenda javës: OK.", inline_html)
        self.assertIn("background-color:#fee2e2;color:#7f1d1d;padding:2px 4px'>Për të premten ose javën tjetër: RREZIK.", inline_html)
        self.assertIn("font-weight:700;color:#7f1d1d", inline_html)
        self.assertIn("font-weight:700;color:#14532d", inline_html)
        self.assertIn("width:42px", inline_html)
        self.assertIn("Klienti kërkoi shtyrjen<br>Rreshti tjetër", inline_html)
        self.assertIn("Përgjigje me ë dhe ç", inline_html)
        self.assertIn("Detyrë e paprekur 119", attached_html)
        self.assertEqual(row.status, "SENT")

    async def test_export_failure_does_not_send_an_incomplete_email(self):
        row = M3ReportingPointsReport(id=uuid.uuid4(), report_date=DAY, manual_answers={},
                                      data={}, status="DRAFT")
        gmail = SimpleNamespace(find_exact=AsyncMock(return_value=None), send_verified=AsyncMock())
        db = AsyncMock()
        with (patch.object(service, "GmailService", return_value=gmail),
              patch("app.services.m3_reporting_points_attachments.report_attachments", side_effect=RuntimeError("Export failed"))):
            with self.assertRaisesRegex(RuntimeError, "Export failed"):
                await service.send_report(db, row, {"to": ["current-m3@example.com"]})
        gmail.send_verified.assert_not_awaited()
        self.assertEqual(row.status, "FAILED")
        self.assertEqual(row.last_error, "Export failed")
        db.commit.assert_awaited_once()

    async def test_api_preserves_other_answers_and_1615_capture_when_saving(self):
        row = M3ReportingPointsReport(id=uuid.uuid4(), report_date=DAY, manual_answers={"underload": "Ruaj kete"},
                                      data={}, status="DRAFT", realization={"percent": 62},
                                      realization_captured_at=datetime.fromisoformat("2026-10-05T16:15:00+02:00"))
        db = SimpleNamespace(get=AsyncMock(return_value=row), refresh=AsyncMock(), commit=AsyncMock(), add=MagicMock())
        with patch.object(api, "locked_report", new=AsyncMock(return_value=row)):
            result = await api.save_answers(row.id, api.ManualAnswersPayload(manual_answers={"ga_reorganization": "Jo"}),
                                             db, SimpleNamespace(id=uuid.uuid4()))
        self.assertEqual(result["manual_answers"], {"underload": "Ruaj kete", "ga_reorganization": "Jo"})
        self.assertEqual(result["realization"], {"percent": 62})
        db.commit.assert_awaited_once()

    async def test_api_manual_send_reads_current_m3_recipients_and_requires_capture(self):
        row = M3ReportingPointsReport(id=uuid.uuid4(), report_date=DAY, manual_answers={}, data={}, status="DRAFT",
                                      generated_at=datetime.now(timezone.utc))
        db = SimpleNamespace(get=AsyncMock(return_value=row), refresh=AsyncMock(), add=MagicMock())
        user = SimpleNamespace(id=uuid.uuid4())
        settings = SimpleNamespace(recipients={"to": ["updated-m3@example.com"], "cc": [], "bcc": []})
        with (patch.object(api, "locked_report", new=AsyncMock(return_value=row)),
              patch.object(api, "get_settings", new=AsyncMock(return_value=settings)),
              patch.object(api, "refresh_report", new=AsyncMock()),
              patch.object(api, "send_report", new=AsyncMock()) as deliver):
            with self.assertRaises(api.HTTPException) as error:
                await api.send(row.id, db, user)
            self.assertEqual(error.exception.status_code, 409)
            deliver.assert_not_awaited()
            row.realization_captured_at = datetime.fromisoformat("2026-10-05T16:15:00+02:00")
            await api.send(row.id, db, user)
        deliver.assert_awaited_once_with(db, row, settings.recipients)


class AccessTests(unittest.TestCase):
    def setUp(self):
        self.row = M3ReportingPointsReport(id=uuid.uuid4(), report_date=datetime.now(service.report_timezone()).date(),
                                          manual_answers={}, data={}, status="DRAFT")
        result = MagicMock()
        result.scalar_one_or_none.return_value = self.row
        result.scalars.return_value.all.return_value = [self.row]
        self.db = SimpleNamespace(execute=AsyncMock(return_value=result), get=AsyncMock(return_value=self.row),
                                  add=MagicMock(), commit=AsyncMock())
        self.user = SimpleNamespace(id=uuid.uuid4(), full_name="Test User", role=UserRole.STAFF)
        self.app = FastAPI()
        self.app.include_router(api.router, prefix="/points")
        self.app.dependency_overrides[api.get_db] = lambda: self.db
        self.app.dependency_overrides[api.get_current_user] = lambda: self.user

    def test_all_roles_can_view_generate_history_and_preview(self):
        with (TestClient(self.app) as client,
              patch.object(api, "locked_report", new=AsyncMock(return_value=self.row)),
              patch.object(api, "refresh_report", new=AsyncMock())):
            for role in UserRole:
                self.user.role = role
                with self.subTest(role=role):
                    for method, url in (("GET", f"/points?report_date={self.row.report_date}"),
                                        ("GET", "/points/history"), ("GET", f"/points/{self.row.id}/preview"),
                                        ("POST", f"/points/generate?report_date={self.row.report_date}")):
                        self.assertEqual(client.request(method, url).status_code, 200, url)

    def test_staff_view_permission_does_not_grant_edit_send_or_recipient_management(self):
        with TestClient(self.app) as client:
            self.assertEqual(client.put(f"/points/{self.row.id}/answers", json={"manual_answers": {"underload": "test"}}).status_code, 403)
            self.assertEqual(client.post(f"/points/{self.row.id}/send").status_code, 403)
            self.assertEqual(client.get("/points/recipients").status_code, 403)
        self.db.commit.assert_not_awaited()

    def test_report_still_requires_authentication(self):
        del self.app.dependency_overrides[api.get_current_user]
        with TestClient(self.app) as client:
            self.assertEqual(client.get("/points/history").status_code, 401)
            self.assertEqual(client.post(f"/points/generate?report_date={self.row.report_date}").status_code, 401)


if __name__ == "__main__":
    unittest.main()
