import io
import uuid
import unittest
from datetime import date, datetime, timezone
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from docx import Document
from fastapi import HTTPException
from PIL import Image
from pydantic import ValidationError
from fastapi import FastAPI
from fastapi.testclient import TestClient

from app.api.routers import m2_reporting_points as api
from app.models.m2_reporting_points import M2ReportingPointsReport
from app.services import m2_reporting_points as service
from app.models.enums import UserRole
from app.services.reporting_points_excel import XLSX_MIME
from openpyxl import load_workbook

DAY = date(2026, 10, 6)


def task(**changes):
    return SimpleNamespace(**{
        "id": uuid.uuid4(), "title": "EF: Detyre", "status": "TODO", "progress_percentage": 0,
        "created_at": DAY, "start_date": DAY, "due_date": DAY,
        "completed_at": None, "is_deadline_important": False, "one_h_marker": None, **changes,
    })


def event(item, field="due_date", before=DAY, after=date(2026, 10, 7)):
    return SimpleNamespace(id=uuid.uuid4(), entity_id=item.id, action=f"task.{field}_changed",
                           created_at=datetime(2026, 10, 6, 8, tzinfo=timezone.utc),
                           before={"value": before.isoformat()}, after={"value": after.isoformat()})


@pytest.mark.parametrize("field", ["start", "due", "start_due"])
def test_original_same_day_task_is_selected_for_each_postponement(field):
    row = task(start_date=date(2026, 10, 7) if "start" in field else DAY,
               due_date=date(2026, 10, 7) if "due" in field else DAY)
    events = [event(row, name) for name in ("start_date", "due_date") if name.split("_")[0] in field]
    selected = service.select_task_sections([row], events, DAY)["postponed"]
    assert len(selected) == 1
    assert selected[0]["postponement_kind"] == field
    assert selected[0]["category"] == "SOT/SOT"


def test_deadline_postponements_do_not_require_creation_today():
    row = task(created_at=date(2026, 9, 1), start_date=date(2026, 9, 15),
               due_date=date(2026, 10, 7), is_deadline_important=True)
    selected = service.select_task_sections([row], [event(row)], DAY)["postponed"]
    assert len(selected) == 1
    assert selected[0]["category"] == "Deadline Important"


def test_regular_old_tasks_and_completion_date_normalization_are_not_postponements():
    old = task(created_at=date(2026, 9, 1), due_date=date(2026, 10, 7))
    done = task(created_at=date(2026, 9, 1), due_date=DAY, completed_at=DAY,
                status="DONE", is_deadline_important=True)
    rows = service.select_task_sections([old, done], [event(old), event(done, before=date(2026, 10, 5), after=DAY)], DAY)
    assert rows["postponed"] == []


@pytest.mark.parametrize("marker", ["M2", "M2_M3"])
def test_all_requested_markers_include_tasks_inside_the_planning_interval(marker):
    rows = [task(one_h_marker=marker), task(one_h_marker=marker, created_at=date(2026, 10, 1),
              start_date=date(2026, 10, 1), due_date=date(2026, 10, 9)),
            task(one_h_marker=marker, status="DONE", completed_at=DAY)]
    selected = service.select_task_sections(rows, [], DAY)["delivery"]
    assert [row["task_id"] for row in selected] == [str(row.id) for row in rows]
    assert selected[0]["marker"] == ("M2/3" if marker == "M2_M3" else marker)


@pytest.mark.parametrize("status", ["TODO", "IN_PROGRESS", "DONE"])
def test_delivery_interval_is_inclusive_and_does_not_depend_on_status(status):
    rows = [task(one_h_marker="M2", status=status, start_date=DAY, due_date=date(2026, 10, 9)),
            task(one_h_marker="M2_M3", status=status, start_date=date(2026, 10, 1), due_date=DAY)]
    assert len(service.select_task_sections(rows, [], DAY)["delivery"]) == 2


def test_m3_only_symbol_is_excluded_even_inside_the_interval():
    assert service.select_task_sections([task(one_h_marker="M3")], [], DAY)["delivery"] == []


def test_overdue_future_missing_or_reversed_intervals_are_excluded():
    rows = [task(one_h_marker="M2", due_date=date(2026, 10, 5), start_date=date(2026, 10, 1)),
            task(one_h_marker="M2", start_date=date(2026, 10, 7), due_date=date(2026, 10, 9)),
            task(one_h_marker="M2", start_date=None), task(one_h_marker="M2", due_date=None),
            task(one_h_marker="M2", start_date=None, due_date=None),
            task(one_h_marker="M2", start_date=date(2026, 10, 9), due_date=date(2026, 10, 1))]
    assert service.select_task_sections(rows, [], DAY)["delivery"] == []


def test_future_tasks_old_completed_tasks_and_unrelated_markers_are_excluded():
    rows = [task(one_h_marker="M2", start_date=date(2026, 10, 7), due_date=date(2026, 10, 7)),
            task(one_h_marker="M3", status="DONE", created_at=date(2026, 10, 5),
                 start_date=date(2026, 10, 5), due_date=date(2026, 10, 5), completed_at=date(2026, 10, 5)),
            task(one_h_marker="FLAG"), task(one_h_marker="M2", start_date=None, due_date=date(2026, 10, 9))]
    assert service.select_task_sections(rows, [], DAY)["delivery"] == []


def report(row):
    return {"report_date": DAY.isoformat(), "manual_answers": {
        "reorganization": "Ndaj detyrat <script>test</script>", f"delivery:{row['task_id']}": "Po, dorëzuar te GA"},
        "data": {"delivery": [row], "postponed": []}}


def complete_export_fixture():
    base = {"assignees": "EF / GA", "department": "DEV", "project": "PrimeFlow",
            "am_pm": "AM", "task_type": "1H", "status": "IN_PROGRESS", "marker": "M2/3",
            "title": "Detyrë me ë dhe ç për kontrollin e dokumenteve", "deadline_important": False}
    postponed = [{**base, "task_id": str(uuid.uuid4()), "postponement_kind": kind,
                  "old_start_date": "2026-10-06", "start_date": "2026-10-07",
                  "old_due_date": "2026-10-06", "due_date": "2026-10-09" if kind == "due" else "2026-10-07",
                  "risk": "RREZIK" if kind == "due" else "OK"} for kind in ("start_due", "due", "start")]
    delivery = [{**base, "task_id": str(uuid.uuid4()), "marker": "M2" if i % 2 else "(M2/3)",
                 "deadline_important": i == 0, "status": "TODO" if i == 2 else "IN_PROGRESS"} for i in range(4)]
    answers = {"reorganization": "Riorganizim i ekipit\nTakim pas pauzës"}
    for row in postponed:
        answers[f"postponed_comment:{row['task_id']}"] = "Koment manual me ë dhe ç\nRreshti i dytë"
    for row, choice in zip(delivery, ("PO", "JO", "", "PO")):
        answers[f"delivery_choice:{row['task_id']}"] = choice
        answers[f"delivery:{row['task_id']}"] = "Përgjigje manuale\n" + "Detajet e dorëzimit duhen ruajtur të plota. " * 4
    return {"report_date": DAY.isoformat(), "manual_answers": answers,
            "data": {"postponed": postponed, "delivery": delivery}}


def test_complete_exports_have_matching_colors_and_complete_content():
    value = complete_export_fixture()
    blocks = service.export_blocks(value)
    tables = [block for block in blocks if block["kind"] == "table"]
    assert len(tables) == 4
    for table in tables[:3]:
        headers = [column[0] for column in table["columns"]]
        assert not {"status", "reason", "comment", "category"}.intersection(headers)
        assert headers[-1] == "manual_comment"
        row = dict(zip(headers, table["rows"][0]))
        assert row["manual_comment"]["fill"] == "#ffffff"
        assert row["risk"]["fill"] == ("#fee2e2" if row["risk"]["text"] == "RREZIK" else "#dcfce7")
    headers = [column[0] for column in tables[-1]["columns"]]
    assert headers[-2:] == ["choice", "answer"]
    for row, fill in zip(tables[-1]["rows"], ("#16a34a", "#dc2626", "#ffffff", "#16a34a")):
        cells = dict(zip(headers, row))
        assert cells["choice"]["fill"] == fill
        assert cells["answer"]["fill"] == "#ffffff"
    html = service.render_html(value)
    assert "TOTALI/DËRGUAR: 4/2" in html
    assert "border-bottom:2px solid #000" in html
    attachments = service.report_attachments(value)
    word = Document(io.BytesIO(attachments[0][1]))
    assert len(word.tables) == 4
    xml = word._element.xml
    assert "w:cantSplit" in xml
    assert not word.styles["Title"]._element.xpath("./w:pPr/w:pBdr")
    for fill in ("16a34a", "dc2626", "dcfce7", "fee2e2"):
        assert f'w:fill="{fill}"' in xml
        assert f"background:{fill if fill.startswith('#') else '#' + fill}" in html
    assert "TOTALI/DËRGUAR: 4/2" in "\n".join(p.text for p in word.paragraphs)
    assert all("Rreshti i dytë" in table.rows[1].cells[-1].text for table in word.tables[:3])
    with Image.open(io.BytesIO(attachments[1][1])) as image:
        colors = {color for _, color in image.getcolors(image.width * image.height)}
        for color in ((22, 163, 74), (220, 38, 38), (220, 252, 231), (254, 226, 226)):
            assert color in colors


def test_preview_and_native_exports_preserve_last_manual_column_and_escape_text():
    item = service.select_task_sections([task(one_h_marker="M2_M3")], [], DAY)["delivery"][0]
    item.update(assignees="EF", department="DEV", project="PrimeFlow", comment="Kontrolluar")
    value = report(item)
    value["manual_answers"][f"delivery_choice:{item['task_id']}"] = "PO"
    value["manual_answers"][f"delivery_choice:{uuid.uuid4()}"] = "PO"
    assert "TOTALI/DËRGUAR: 1/1" in service.render_plain_text(value)
    assert service.table_columns("delivery")[-1][0] == "answer"
    html = service.render_html(value)
    assert "<script>test</script>" not in html
    assert "&lt;script&gt;test&lt;/script&gt;" in html
    assert "Po, dorëzuar te GA" in html
    assert "M2/3" in service.render_plain_text(value)
    attachments = service.report_attachments(value)
    word = Document(io.BytesIO(attachments[0][1]))
    table = word.tables[-1]
    headings = [cell.text for cell in table.rows[0].cells]
    assert not {"STATUSI", "ARSYEJA", "KOMENTI"}.intersection(headings)
    assert len(headings) == 10
    assert table.rows[1].cells[headings.index("PO/JO")].text == "PO"
    marker_index = headings.index("SIMBOLI")
    assert headings[marker_index + 1] == "TITULLI"
    assert table.rows[1].cells[marker_index].text == "M2/3"
    assert table.rows[1].cells[marker_index + 1].text == item["title"]
    assert "SIMBOLI: M2/3 | TITULLI: EF: Detyre" in service.render_plain_text(value)
    assert table.rows[0].cells[-1].text == "PERGJIGJJA MANUALE"
    assert table.rows[1].cells[-1].text == "Po, dorëzuar te GA"
    with Image.open(io.BytesIO(attachments[1][1])) as image:
        assert image.width > 1000
        assert image.height > 100


@pytest.mark.parametrize("key", ["overload", "delivery:invalid", "other:key"])
def test_answer_payload_rejects_unknown_keys(key):
    with pytest.raises(ValidationError):
        api.ManualAnswersPayload(manual_answers={key: "Test"})


@pytest.mark.parametrize("choice", ["", "PO", "JO"])
def test_delivery_dropdown_accepts_supported_choices(choice):
    key = f"delivery_choice:{uuid.uuid4()}"
    assert api.ManualAnswersPayload(manual_answers={key: choice}).manual_answers[key] == choice
    with pytest.raises(ValidationError):
        api.ManualAnswersPayload(manual_answers={key: "OTHER"})


class ReportPersistenceTests(unittest.IsolatedAsyncioTestCase):
    async def test_unfinished_priority_reuses_after_break_rows_and_preserves_existing_sections(self):
        rows = [
            ["1", "ER", "GD", "AM", "TODO", "DUE SOT", "PRJK", "ER: Detyre", "SOT"],
            ["2", "EF", "DEV", "AM/PM", "IN_PROGRESS", "DEADLINE / 08:00", "1H", "08:00 EF: Test <task>", "SOT"],
        ]
        db = AsyncMock()
        with patch.object(service.m3, "build_task_data", new=AsyncMock(return_value={"postponed": [], "delivery": []})), \
             patch.object(service, "build_unfinished_priority_task_rows", new=AsyncMock(return_value=rows)) as shared_rows:
            data = await service.build_task_data(db, DAY)
        shared_rows.assert_awaited_once_with(db, DAY)
        assert data["delivery"] == data["postponed"] == []
        assert [row["title"] for row in data["unfinished_priority"]] == [row[7] for row in rows]
        assert data["unfinished_priority"][0]["deadline_important"] is False
        priority = data["unfinished_priority"][1]
        assert priority["eight_am"] and priority["deadline_important"]
        assert priority["am_pm"] == "AM/PM" and priority["due_label"] == "SOT"
        value = {"report_date": DAY.isoformat(), "manual_answers": {}, "data": data}
        blocks = service.export_blocks(value)
        titles = [block["text"] for block in blocks if block["kind"] == "text" and block["level"] == 2]
        assert titles == [service.MANUAL_POINTS["reorganization"], *service.AUTO_TITLES.values()]
        table = next(block for block in blocks if block["kind"] == "table")
        assert [column[1] for column in table["columns"]] == ["NR", "KUSH", "DEP", "AM/PM", "LLOJI", "TIPI", "TITULLI", "DUE DATE"]
        assert all(cell["fill"] == "#dc2626" and cell["eight_am"] for cell in table["rows"][1])
        html = service.render_html(value)
        assert "08:00 EF: Test &lt;task&gt;" in html
        assert "DUE DATE: SOT" in service.render_plain_text(value)
        attachments = service.report_attachments(value)
        word = Document(io.BytesIO(attachments[0][1]))
        assert word.tables[0].rows[2].cells[6].text == priority["title"]
        assert word.tables[0].rows[2].cells[7].text == "SOT"

    async def test_email_preparation_uses_full_html_word_png_and_excel_without_real_email(self):
        fixture = complete_export_fixture()
        row = M2ReportingPointsReport(id=uuid.uuid4(), report_date=DAY, manual_answers=fixture["manual_answers"], data=fixture["data"], status="DRAFT")
        gmail = SimpleNamespace(find_exact=AsyncMock(return_value=None), send_verified=AsyncMock(return_value={"id": "mock-message"}))
        db = AsyncMock()
        with patch.object(service, "GmailService", return_value=gmail):
            await service.send_report(db, row, {"to": ["test@example.com"], "cc": [], "bcc": []})
        args = gmail.send_verified.call_args.args
        attachments = gmail.send_verified.call_args.kwargs["attachments"]
        assert "TOTALI/DËRGUAR: 4/2" in args[2] and "TOTALI/DËRGUAR: 4/2" in args[3]
        assert {name.rsplit(".", 1)[-1] for name, _, _ in attachments} == {"html", "docx", "png", "xlsx"}
        name, content, mime = next(item for item in attachments if item[0].endswith(".xlsx"))
        assert name == f"pikat_m2_{DAY.isoformat()}.xlsx" and mime == XLSX_MIME
        sheet = load_workbook(io.BytesIO(content)).active
        assert "TOTALI/DËRGUAR: 4/2" in [cell.value for row in sheet for cell in row]
        assert next(content for name, content, _ in attachments if name.endswith(".html")).decode("utf-8") == args[3]
        assert row.status == "SENT" and row.gmail_message_id == "mock-message"
        db.commit.assert_awaited_once()
    async def test_postponement_comments_save_and_export_without_affecting_delivery_answers(self):
        item = task(due_date=date(2026, 10, 7))
        data = service.select_task_sections([item], [event(item)], DAY)
        key = f"postponed_comment:{item.id}"
        row = M2ReportingPointsReport(id=uuid.uuid4(), report_date=DAY, manual_answers={"reorganization": "Po"}, data=data, status="DRAFT")
        with patch.object(api, "_by_id", new=AsyncMock(return_value=row)), patch.object(api, "locked_report", new=AsyncMock(return_value=row)), patch.object(api, "add_audit_log"):
            result = await api.save_answers(row.id, api.ManualAnswersPayload(manual_answers={key: "Koment manual <test>"}), AsyncMock(), SimpleNamespace(id=uuid.uuid4()))
        assert result["manual_answers"] == {"reorganization": "Po", key: "Koment manual <test>"}
        assert "Koment manual &lt;test&gt;" in service.render_html(result)
        attachments = service.report_attachments(result)
        table = Document(io.BytesIO(attachments[0][1])).tables[0]
        assert table.rows[0].cells[-1].text == "KOMENT MANUAL"
        assert table.rows[1].cells[-1].text == "Koment manual <test>"

    async def test_refresh_preserves_reorganization_and_per_task_answers(self):
        row = M2ReportingPointsReport(id=uuid.uuid4(), report_date=DAY, manual_answers={"reorganization": "Po"}, data={}, status="DRAFT")
        with patch.object(service, "datetime") as clock, patch.object(service, "build_task_data", new=AsyncMock(return_value={"delivery": []})):
            clock.now.return_value = datetime(2026, 10, 6, 10, tzinfo=timezone.utc)
            await service.refresh_report(AsyncMock(), row)
        assert row.manual_answers == {"reorganization": "Po"}
        assert row.data == {"delivery": []}


    async def test_saving_answers_merges_and_rejects_task_from_another_report(self):
        task_id = str(uuid.uuid4())
        row = M2ReportingPointsReport(id=uuid.uuid4(), report_date=DAY, manual_answers={"reorganization": "Po"},
                                      data={"delivery": [{"task_id": task_id}]}, status="DRAFT")
        db = AsyncMock()
        with patch.object(api, "_by_id", new=AsyncMock(return_value=row)), patch.object(api, "locked_report", new=AsyncMock(return_value=row)), patch.object(api, "add_audit_log"):
            result = await api.save_answers(row.id, api.ManualAnswersPayload(manual_answers={f"delivery:{task_id}": "Dorëzuar"}), db, SimpleNamespace(id=uuid.uuid4()))
            assert result["manual_answers"] == {"reorganization": "Po", f"delivery:{task_id}": "Dorëzuar"}
            with pytest.raises(HTTPException) as error:
                await api.save_answers(row.id, api.ManualAnswersPayload(manual_answers={f"delivery:{uuid.uuid4()}": "Jo"}), db, SimpleNamespace(id=uuid.uuid4()))
            assert error.value.status_code == 422
            row.status = "SENT"
            row.sent_at = datetime(2026, 10, 6, 10, tzinfo=timezone.utc)
            result = await api.save_answers(row.id, api.ManualAnswersPayload(manual_answers={"reorganization": "Ndrysho"}), db, SimpleNamespace(id=uuid.uuid4()))
            assert result["manual_answers"]["reorganization"] == "Ndrysho"
            assert result["status"] == "DRAFT"
            assert result["sent_at"] == row.sent_at.isoformat()


    async def test_manual_send_delivers_every_time_with_latest_answers(self):
        row = M2ReportingPointsReport(id=uuid.uuid4(), report_date=DAY, manual_answers={"reorganization": "Fillimi"},
                                      data={}, status="SENT", gmail_message_id="previous-message")
        gmail = SimpleNamespace(find_exact=AsyncMock(return_value={"id": "previous-message"}),
                                send_verified=AsyncMock(side_effect=[{"id": f"new-{i}"} for i in range(3)]))
        db = AsyncMock()
        with patch.object(service, "GmailService", return_value=gmail), patch.object(service, "report_attachments", return_value=[]):
            await service.send_report(db, row, {"to": ["test@example.com"]})
            first_sent_at = row.sent_at
            row.manual_answers = {"reorganization": "Ndryshimi"}
            await service.send_report(db, row, {"to": ["test@example.com"]})
            await service.send_report(db, row, {"to": ["test@example.com"]})
        assert gmail.send_verified.await_count == 3
        gmail.find_exact.assert_not_awaited()
        assert "Fillimi" in gmail.send_verified.call_args_list[0].args[2]
        assert all("Ndryshimi" in call.args[2] for call in gmail.send_verified.call_args_list[1:])
        assert row.status == "SENT" and row.gmail_message_id == "new-2"
        assert row.sent_at >= first_sent_at
        assert db.commit.await_count == 3

    async def test_sent_report_can_be_regenerated_without_losing_answers_or_last_send(self):
        sent_at = datetime(2026, 10, 6, 9, tzinfo=timezone.utc)
        row = M2ReportingPointsReport(id=uuid.uuid4(), report_date=DAY, manual_answers={"reorganization": "Po"},
                                      data={}, status="SENT", sent_at=sent_at, gmail_message_id="old-message")
        db = AsyncMock()
        with patch.object(api, "datetime") as clock, patch.object(service, "datetime") as service_clock, \
             patch.object(api, "locked_report", new=AsyncMock(return_value=row)), patch.object(api, "add_audit_log"), \
             patch.object(service, "build_task_data", new=AsyncMock(return_value={"unfinished_priority": [{"title": "Detyre e re"}]})):
            clock.now.return_value = service_clock.now.return_value = datetime(2026, 10, 6, 11, tzinfo=timezone.utc)
            result = await api.generate(DAY, db, SimpleNamespace(id=uuid.uuid4()))
        assert result["status"] == "DRAFT"
        assert result["data"]["unfinished_priority"][0]["title"] == "Detyre e re"
        assert result["manual_answers"] == {"reorganization": "Po"}
        assert row.sent_at == sent_at and row.gmail_message_id == "old-message"
        db.commit.assert_awaited_once()

    async def test_send_route_refreshes_and_audits_every_resend(self):
        row = M2ReportingPointsReport(id=uuid.uuid4(), report_date=DAY, manual_answers={}, data={},
                                      status="SENT", generated_at=datetime(2026, 10, 6, 9, tzinfo=timezone.utc))
        with patch.object(api, "datetime") as clock, patch.object(api, "_by_id", new=AsyncMock(return_value=row)), \
             patch.object(api, "locked_report", new=AsyncMock(return_value=row)), \
             patch.object(api, "get_settings", new=AsyncMock(return_value=SimpleNamespace(recipients={"to": ["test@example.com"]}))), \
             patch.object(api, "refresh_report", new=AsyncMock()) as refresh, \
             patch.object(api, "send_report", new=AsyncMock()) as send, patch.object(api, "add_audit_log") as audit:
            clock.now.return_value = datetime(2026, 10, 6, 11, tzinfo=timezone.utc)
            for _ in range(3):
                await api.send(row.id, AsyncMock(), SimpleNamespace(id=uuid.uuid4()))
        assert refresh.await_count == send.await_count == audit.call_count == 3

    async def test_failed_resend_keeps_last_successful_delivery_and_can_be_retried(self):
        sent_at = datetime(2026, 10, 6, 9, tzinfo=timezone.utc)
        row = M2ReportingPointsReport(id=uuid.uuid4(), report_date=DAY, manual_answers={}, data={},
                                      status="SENT", sent_at=sent_at, gmail_message_id="previous-message")
        gmail = SimpleNamespace(send_verified=AsyncMock(side_effect=[RuntimeError("SMTP unavailable"), {"id": "retry-message"}]))
        db = AsyncMock()
        with patch.object(service, "GmailService", return_value=gmail), patch.object(service, "report_attachments", return_value=[]):
            with pytest.raises(RuntimeError, match="SMTP unavailable"):
                await service.send_report(db, row, {"to": ["test@example.com"]})
            assert row.status == "FAILED" and row.sent_at == sent_at and row.gmail_message_id == "previous-message"
            await service.send_report(db, row, {"to": ["test@example.com"]})
        assert row.status == "SENT" and row.gmail_message_id == "retry-message" and row.last_error is None


    async def test_m2_recipient_configuration_is_used(self):
        db = AsyncMock()
        result = MagicMock()
        result.scalars.return_value.first.return_value = SimpleNamespace(manual_recipients={"to": ["m2@example.com"]})
        db.execute.return_value = result
        response = await api.recipients(db, SimpleNamespace())
        assert response["recipients"]["to"] == ["m2@example.com"]
        assert "reporting_points_settings" in str(db.execute.call_args.args[0])


def test_staff_can_read_but_cannot_edit_send_or_read_recipients():
    app = FastAPI()
    app.include_router(api.router, prefix="/points")
    app.dependency_overrides[api.get_current_user] = lambda: SimpleNamespace(role=UserRole.STAFF, full_name="Staff")
    db = AsyncMock()
    result = MagicMock()
    result.scalars.return_value.all.return_value = []
    db.execute.return_value = result
    app.dependency_overrides[api.get_db] = lambda: db
    with TestClient(app) as client:
        assert client.get("/points/history").status_code == 200
        assert client.get("/points/recipients").status_code == 403
        assert client.put(f"/points/{uuid.uuid4()}/answers", json={"manual_answers": {"reorganization": "Po"}}).status_code == 403
        assert client.post(f"/points/{uuid.uuid4()}/send").status_code == 403
        del app.dependency_overrides[api.get_current_user]
        assert client.get("/points/history").status_code == 401
    db.commit.assert_not_awaited()


def test_migration_creates_separate_m2_storage_without_changing_m3():
    from alembic.migration import MigrationContext
    from alembic.operations import Operations
    from io import StringIO
    from tests.test_migration_graph import migration_scripts

    script = migration_scripts().get_revision("0143_m2_reporting_points")
    buffer = StringIO()
    context = MigrationContext.configure(dialect_name="postgresql", opts={"as_sql": True, "output_buffer": buffer})
    with patch.object(script.module, "op", Operations(context)):
        script.module.upgrade()
    sql = buffer.getvalue()
    assert "CREATE TABLE m2_reporting_points_reports" in sql
    assert "manual_answers JSONB" in sql
    assert "UNIQUE (report_date)" in sql
    assert "m3_reporting_points_reports" not in sql
