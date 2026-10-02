from __future__ import annotations

import io
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from docx import Document
from PIL import Image, ImageDraw

from app.services import m3_reporting_points as report_service
from app.services import m3_reporting_points_attachments as exports


def sample_report():
    base = {"assignees": "EF", "department": "DEV", "project": "PRJ", "am_pm": "PM", "task_type": "SYS",
            "created_date": "2026-10-02", "old_start_date": "2026-10-02", "old_due_date": "2026-10-02",
            "start_date": "2026-10-05", "due_date": "2026-10-09", "postponement_kind": "start_due",
            "reason": "Në pritje të klientit", "comment": "Klienti kërkoi shtyrjen\nDuhet informacion tjetër", "risk": "RREZIK"}
    moved = [{**base, "task_id": str(i), "title": f"Detyrë me status {status}", "status": status, "risk": "OK" if i == 0 else "RREZIK"}
             for i, status in enumerate(report_service.STATUS_COLORS)]
    moved[0]["marker"] = "👁"
    moved[1]["marker"] = "!"
    moved += [{**base, "task_id": "due", "title": "Shtyrje vetëm DUE", "status": "TODO", "postponement_kind": "due", "deadline_important": True},
              {**base, "task_id": "start", "title": "08:00 Shtyrje vetëm START", "status": "IN_PROGRESS", "postponement_kind": "start"}]
    return {"report_date": "2026-10-02", "manual_answers": {key: f"Përgjigje {key} me ë dhe ç\nRreshti i dytë" for key in report_service.MANUAL_POINTS},
            "data": {"postponed": moved, "untouched": [{**base, "title": "SYS TODO sot", "status": "TODO"}, {**base, "title": "TODO projekt", "task_type": "1H", "status": "TODO"}],
                     "same_day": [{**base, "title": "DONE në fund", "status": "DONE"}, {**base, "title": "Aktive përpara DONE", "status": "IN_PROGRESS"}],
                     "ga_postponed": [{**base, "title": "DEADLINE IMPORTANT për GA", "status": "WAITING_CLIENT", "category": "Deadline Important"}]},
            "realization": {"percent": 62, "comment": "Jemi mbi 50%", "departments": [
                {"department_id": "dev", "code": "DEV", "name": "Development", "percent": 75, "employees": 5, "comment": "Jemi mbi 50%"},
                {"department_id": "gd", "code": "GD", "name": "Graphic Design", "percent": 25, "employees": 2, "comment": "Jemi nen 50%"}]},
            "realization_captured_at": "2026-10-02T16:15:00+02:00"}


class ReportingPointsAttachmentTests(unittest.TestCase):
    def test_word_contains_all_sections_editable_tables_colors_and_repeating_headers(self):
        report = sample_report()
        data = exports.render_docx(report)
        document = Document(io.BytesIO(data))
        paragraphs = "\n".join(p.text for p in document.paragraphs)
        for title in (report_service.M3_TITLE, report_service.GA_TITLE, *report_service.MANUAL_POINTS.values(), *report_service.AUTO_TITLES.values()):
            self.assertIn(title, paragraphs)
        for answer in report["manual_answers"].values():
            self.assertIn(answer, paragraphs)
        self.assertIn("62%", paragraphs)
        self.assertIn("16:15", paragraphs)
        table_text = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
        for rows in report["data"].values():
            for row in rows:
                if row["title"] == "DONE në fund":
                    self.assertNotIn(row["title"], table_text)
                    continue
                self.assertIn(row["title"], table_text)
                self.assertIn(row["reason"], table_text)
                self.assertIn(row["comment"], table_text)
        self.assertIn("Aktive përpara DONE", table_text)
        self.assertNotIn("DONE në fund", table_text)
        self.assertEqual(len(document.tables), 8)
        self.assertIn("DETYRAT E SISTEMIT PA PROGRES:", paragraphs)
        self.assertIn("DET FT DHE PRJK PA PROGRES:", paragraphs)
        self.assertIn("75%", table_text)
        self.assertIn("25%", table_text)
        self.assertIn("👁 Detyrë me status TODO", table_text)
        self.assertIn("! Detyrë me status IN_PROGRESS", table_text)
        self.assertIn("START: 02.10.2026\nDUE: 02.10.2026", table_text)
        self.assertNotIn("STATUSI", table_text)
        self.assertNotIn("CREATION / START / DUE", table_text)
        self.assertNotIn("Progresi %", table_text)
        with ZipFile(io.BytesIO(data)) as archive:
            xml = archive.read("word/document.xml").decode().lower()
        for color in (*report_service.STATUS_COLORS.values(), exports.GREEN, exports.RED, "#dc2626"):
            self.assertIn(f'w:fill="{color.lstrip("#").lower()}"', xml)
        self.assertIn('w:sz="18"', xml)
        self.assertIn('w:color="000000"', xml)
        self.assertEqual(xml.count("<w:tblheader"), 8)
        self.assertIn('w:color="dc2626"', xml)
        page = document.sections[0]
        self.assertGreater(page.page_width, page.page_height)
        for table in document.tables:
            self.assertLessEqual(sum(column.width for column in table.columns), page.page_width - page.left_margin - page.right_margin)

    def test_png_renders_complete_report_without_clipping_long_text_and_keeps_colors(self):
        report = sample_report()
        report["manual_answers"]["underload"] = "Përgjigje e gjatë " * 70 + "FUNDI MANUAL"
        report["data"]["ga_postponed"][0]["comment"] = "X" * 150 + " FUND KOMENTI"
        calls = []
        original = ImageDraw.ImageDraw.text

        def record(draw, xy, text, *args, **kwargs):
            calls.append((xy, text, kwargs["font"]))
            return original(draw, xy, text, *args, **kwargs)

        with patch.object(ImageDraw.ImageDraw, "text", record):
            data = exports.render_png(report)
        image = Image.open(io.BytesIO(data))
        image.load()
        self.assertGreater(image.width, 1400)
        self.assertGreater(image.height, 1800)
        visible = "".join(value for _, value, _ in calls)
        self.assertNotIn("CREATION / START / DUE", visible)
        self.assertNotIn("DONE në fund", visible)
        for title in (*report_service.AUTO_TITLES.values(), report_service.GA_TITLE, "FUNDI MANUAL", "FUND KOMENTI", "62%", "75%", "25%", "DET FT DHE PRJK PA PROGRES:", "DETYRAT E SISTEMIT PA PROGRES:"):
            self.assertIn(title.replace(" ", ""), visible.replace(" ", ""))
        self.assertIn("X" * 150, visible.replace(" ", ""))
        probe = ImageDraw.Draw(image)
        for xy, value, font in calls:
            bounds = probe.textbbox(xy, value, font=font)
            self.assertGreaterEqual(bounds[0], 0)
            self.assertGreaterEqual(bounds[1], 0)
            self.assertLessEqual(bounds[2], image.width - 30)
            self.assertLessEqual(bounds[3], image.height - 30)
        colors = {color for _, color in image.getcolors(maxcolors=image.width * image.height)}
        for color in (*report_service.STATUS_COLORS.values(), exports.GREEN, exports.RED, "#dc2626"):
            self.assertIn(tuple(bytes.fromhex(color.lstrip("#"))), colors)
        self.assertIn((0, 0, 0), colors)

    def test_attachment_formats_and_names(self):
        attachments = exports.report_attachments(sample_report())
        self.assertEqual([item[2] for item in attachments], [exports.DOCX_MIME, "image/png"])
        self.assertEqual([item[0] for item in attachments], ["PrimeFlow-PIKAT-M3-GA-2026-10-02.docx", "PrimeFlow-PIKAT-M3-GA-2026-10-02.png"])
        self.assertTrue(attachments[0][1].startswith(b"PK"))
        self.assertTrue(attachments[1][1].startswith(b"\x89PNG\r\n\x1a\n"))
