import unittest
import io
import uuid
from datetime import date, datetime, timezone
from types import SimpleNamespace

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from starlette.datastructures import UploadFile

from app.api.routers.exports import (
    OPEN_TASK_BASELINE_STATUS_VALUES,
    OPEN_TASK_EXPORT_HEADERS,
    OPEN_TASK_WHEN_VALUES,
    _normalize_open_task_baseline_value,
    _open_task_baseline_column_numbers,
    _open_task_difference,
    _open_task_planned_status,
    _open_task_planning_when,
    _open_task_source_label,
    _open_task_wrapped_line_count,
    export_open_tasks_xlsx,
    import_open_tasks_baseline,
)
from app.models.enums import UserRole


class TestOpenTasksExportLayout(unittest.TestCase):
    def test_wrapped_line_count_accounts_for_column_width(self) -> None:
        title = "AT/OH:EF/RA: ASC: DEF KO1/KO2/KOF PER KUZHINA (CLAIMS)"

        self.assertEqual(_open_task_wrapped_line_count(title, 44), 2)

    def test_wrapped_line_count_accounts_for_explicit_newlines(self) -> None:
        title = "First line\nSecond line\nThird line"

        self.assertEqual(_open_task_wrapped_line_count(title, 44), 3)

    def test_px_jav_task_uses_ga_ka_source_group(self) -> None:
        task = SimpleNamespace(
            ga_note_origin_id=None,
            plan_note_origin_id="plan-note-id",
            system_template_origin_id=None,
            project_id="project-id",
            is_bllok=False,
            is_r1=False,
            is_1h_report=False,
            is_personal=False,
        )

        self.assertEqual(_open_task_source_label(task), "GA/KA")

    def test_comparison_columns_are_paired_in_requested_order(self) -> None:
        self.assertEqual(
            OPEN_TASK_EXPORT_HEADERS[16:21],
            ["WHEN", "WHEN PLANNED", "STATUS MANUAL", "STATUS PLANNED", "KOMENT"],
        )

    def test_baseline_columns_are_found_in_legacy_and_current_layouts(self) -> None:
        legacy_headers = [
            "STATUS" if header == "STATUS MANUAL" else header
            for header in OPEN_TASK_EXPORT_HEADERS
            if header != "TASK ID"
        ]
        for headers, expected in (
            (legacy_headers, (16, 18, 20)),
            (OPEN_TASK_EXPORT_HEADERS, (17, 19, 21)),
        ):
            workbook = Workbook()
            worksheet = workbook.active
            for column, header in enumerate(headers, start=1):
                worksheet.cell(4, column, header)
            self.assertEqual(_open_task_baseline_column_numbers(worksheet), expected)

    def test_planning_when_uses_upcoming_week_as_this_week(self) -> None:
        current_monday = date(2026, 9, 7)

        self.assertEqual(_open_task_planning_when(date(2026, 9, 11), current_monday), "SOT")
        self.assertEqual(_open_task_planning_when(date(2026, 9, 14), current_monday), "THIS WEEK")
        self.assertEqual(_open_task_planning_when(date(2026, 9, 18), current_monday), "THIS WEEK")
        self.assertEqual(_open_task_planning_when(date(2026, 9, 21), current_monday), "NEXT WEEK")
        self.assertEqual(_open_task_planning_when(date(2026, 9, 25), current_monday), "NEXT WEEK")
        self.assertEqual(_open_task_planning_when(date(2026, 9, 28), current_monday), "FUTURE")

    def test_planning_when_accepts_datetime_and_leaves_unplanned_dates_blank(self) -> None:
        current_monday = date(2026, 9, 7)

        self.assertEqual(
            _open_task_planning_when(datetime(2026, 9, 14, 8, 0, tzinfo=timezone.utc), current_monday),
            "THIS WEEK",
        )
        self.assertEqual(_open_task_planning_when(date(2026, 9, 10), current_monday), "")
        self.assertEqual(_open_task_planning_when(None, current_monday), "")

    def test_planned_status_uses_primeflow_task_flags(self) -> None:
        def task(**overrides):
            values = {"is_1h_report": False, "is_personal": False, "is_bllok": False, "is_r1": False}
            values.update(overrides)
            return SimpleNamespace(**values)

        self.assertEqual(_open_task_planned_status(task(is_1h_report=True)), "1H")
        self.assertEqual(_open_task_planned_status(task(is_personal=True)), "PERSONAL")
        self.assertEqual(_open_task_planned_status(task(is_bllok=True)), "BLLOK")
        self.assertEqual(_open_task_planned_status(task(is_r1=True)), "R1")
        self.assertEqual(_open_task_planned_status(task()), "")

    def test_only_non_matching_planned_values_are_displayed(self) -> None:
        self.assertEqual(_open_task_difference("THIS WEEK", "THIS WEEK"), "")
        self.assertEqual(_open_task_difference("THIS WEEK", "NEXT WEEK"), "NEXT WEEK")
        self.assertEqual(_open_task_difference("1H", "1H"), "")
        self.assertEqual(_open_task_difference("1H", "BLLOK"), "BLLOK")
        self.assertEqual(_open_task_difference("THIS WEEK", None), "UNPLANNED")
        self.assertEqual(_open_task_difference("1H", ""), "UNPLANNED")
        self.assertEqual(_open_task_difference(None, None), "")

    def test_baseline_import_normalizes_and_validates_values(self) -> None:
        self.assertEqual(
            _normalize_open_task_baseline_value(" this   week ", OPEN_TASK_WHEN_VALUES, "WHEN"),
            "THIS WEEK",
        )
        self.assertEqual(
            _normalize_open_task_baseline_value("bllok", OPEN_TASK_BASELINE_STATUS_VALUES, "STATUS"),
            "BLLOK",
        )
        self.assertIsNone(_normalize_open_task_baseline_value("", OPEN_TASK_WHEN_VALUES, "WHEN"))
        with self.assertRaises(HTTPException):
            _normalize_open_task_baseline_value("THIS MONTH", OPEN_TASK_WHEN_VALUES, "WHEN")


class _FakeResult:
    def __init__(self, rows):
        self.rows = rows

    def scalars(self):
        return self

    def unique(self):
        return self

    def all(self):
        return self.rows


class _FakeSession:
    def __init__(self, results):
        self.results = iter(results)
        self.added = []
        self.committed = False

    async def execute(self, _statement):
        return _FakeResult(next(self.results))

    def add(self, value):
        self.added.append(value)

    async def commit(self):
        self.committed = True


class TestOpenTasksExportWorkbook(unittest.IsolatedAsyncioTestCase):
    async def test_import_saves_manual_values_for_the_planning_week(self) -> None:
        task_id = uuid.uuid4()
        workbook = Workbook()
        ws = workbook.active
        ws.title = "OPEN TASKS"
        legacy_headers = [
            "STATUS" if header == "STATUS MANUAL" else header
            for header in OPEN_TASK_EXPORT_HEADERS
            if header != "TASK ID"
        ]
        for column, header in enumerate(legacy_headers, start=1):
            ws.cell(4, column, header)
        ws.cell(5, 16, " this week ")
        ws.cell(5, 18, "bllok")
        ws.cell(5, 20, "Move after customer reply")
        metadata = workbook.create_sheet("_PRIMEFLOW")
        metadata.sheet_state = "veryHidden"
        metadata["A1"] = "OPEN_TASKS_BASELINE_V1"
        metadata["B1"] = "2026-09-14"
        metadata["A2"] = "EXCEL ROW"
        metadata["B2"] = "TASK ID"
        metadata["A3"] = 5
        metadata["B3"] = str(task_id)
        content = io.BytesIO()
        workbook.save(content)
        content.seek(0)

        task = SimpleNamespace(id=task_id, department_id=None)
        db = _FakeSession([[task], []])
        user = SimpleNamespace(
            id=uuid.uuid4(),
            role=UserRole.ADMIN,
            department_id=None,
        )
        result = await import_open_tasks_baseline(
            file=UploadFile(filename="OPEN_TASKS.xlsx", file=content),
            db=db,
            user=user,
        )

        self.assertEqual(result, {"planning_week_start": "2026-09-14", "imported": 1})
        self.assertTrue(db.committed)
        self.assertEqual(len(db.added), 1)
        self.assertEqual(db.added[0].when_value, "THIS WEEK")
        self.assertEqual(db.added[0].status_value, "BLLOK")
        self.assertEqual(db.added[0].comment_value, "Move after customer reply")

    async def test_workbook_contains_manual_and_planned_pairs_with_metadata(self) -> None:
        task_id = uuid.uuid4()
        task = SimpleNamespace(
            id=task_id,
            title="Plan customer report",
            description=None,
            project_id=None,
            department_id=None,
            assigned_to=None,
            assignees=[],
            ga_note_origin_id=None,
            plan_note_origin_id=None,
            system_template_origin_id=None,
            status="TODO",
            priority="NORMAL",
            finish_period="AM",
            start_date=datetime(2026, 9, 7, tzinfo=timezone.utc),
            due_date=datetime(2026, 9, 21, tzinfo=timezone.utc),
            created_at=datetime(2026, 9, 1, tzinfo=timezone.utc),
            is_bllok=False,
            is_1h_report=True,
            is_r1=False,
            is_personal=False,
        )
        baseline = SimpleNamespace(
            task_id=task_id,
            when_value="THIS WEEK",
            status_value="1H",
            comment_value="Before planning",
        )
        db = _FakeSession([[task], [], [baseline]])
        user = SimpleNamespace(
            id=uuid.uuid4(),
            role=UserRole.ADMIN,
            department_id=None,
            full_name="Test User",
            username="tester",
        )

        response = await export_open_tasks_xlsx(
            department_id=None,
            user_id=None,
            filter="all",
            status_filter="all",
            type_filter=None,
            search=None,
            this_week_start=date(2026, 9, 7),
            db=db,
            user=user,
        )
        content = b"".join([chunk async for chunk in response.body_iterator])
        workbook = load_workbook(io.BytesIO(content), data_only=False)
        ws = workbook["OPEN TASKS"]

        self.assertEqual([ws.cell(4, column).value for column in range(17, 22)], OPEN_TASK_EXPORT_HEADERS[16:21])
        self.assertEqual(ws.cell(4, 2).value, "TASK ID")
        self.assertEqual(ws.cell(5, 2).value, str(task_id))
        self.assertEqual(
            [ws.cell(5, column).value for column in range(17, 22)],
            ["THIS WEEK", "NEXT WEEK", "1H", None, "Before planning"],
        )
        self.assertEqual(ws.freeze_panes, "C5")
        self.assertEqual(workbook["_PRIMEFLOW"].sheet_state, "veryHidden")
        self.assertEqual(workbook["_PRIMEFLOW"]["B1"].value, "2026-09-14")
        self.assertEqual(workbook["_PRIMEFLOW"]["B3"].value, str(task_id))


if __name__ == "__main__":
    unittest.main()
