import io
from collections import Counter
from unittest.mock import AsyncMock, patch

import pytest
from openpyxl import load_workbook

from app.services import m2_reporting_points as m2
from app.services import m3_reporting_points_attachments as m3
from app.services import reporting_points_excel as excel
from tests.test_m2_reporting_points import complete_export_fixture
from tests.test_m3_reporting_points_attachments import sample_report


def load(blocks, name="Report"):
    return load_workbook(io.BytesIO(excel.render_xlsx(blocks, sheet_name=name)), rich_text=True).active


def populated(sheet):
    return [cell for row in sheet for cell in row if cell.value is not None]


@pytest.mark.parametrize("report,export,name", [
    (complete_export_fixture, m2.export_blocks, "PIKAT M2"),
    (sample_report, m3.export_blocks, "PIKAT M3 dhe GA"),
])
def test_excel_preserves_every_view_value_group_color_and_manual_answer(report, export, name):
    blocks = export(report())
    sheet = load(blocks, name)
    expected = []
    expected_fills = set()
    for block in blocks:
        if block["kind"] == "text":
            expected.append(block["text"])
        elif block["kind"] == "legend":
            expected.extend(text.strip() for text, _, _ in block["runs"])
        else:
            expected.extend(label for _, label, _ in block["columns"])
            for row in block["rows"]:
                for (key, _, _), cell in zip(block["columns"], row):
                    value = cell["text"]
                    if key == "percent" and value.endswith("%"):
                        expected.append(str(float(value[:-1]) / 100))
                    else:
                        expected.extend(value.split("\n", 1) if cell.get("divider") else [value])
                    expected_fills.add(cell["fill"].lstrip("#").upper())
    cells = populated(sheet)
    assert Counter(str(cell.value) for cell in cells) == Counter(expected)
    assert expected_fills <= {cell.fill.fgColor.rgb[-6:].upper() for cell in cells}
    assert sheet.title == name and sheet.sheet_view.showGridLines is False
    assert all(cell.alignment.wrap_text for cell in cells)
    assert sheet.page_setup.orientation == "landscape"
    assert sheet.page_setup.fitToWidth == 1 and sheet.page_setup.fitToHeight == 0
    assert all(height.height <= 400 for height in sheet.row_dimensions.values())


def test_stacked_dates_have_black_dividers_and_eight_am_rows_keep_red_outline():
    report = sample_report()
    report["data"]["postponed"][0]["eight_am"] = True
    sheet = load(m3.export_blocks(report))
    cells = populated(sheet)
    start = next(cell for cell in cells if cell.value == "START: 02.10.2026")
    due = next(cell for cell in cells if cell.value == "DUE: 02.10.2026")
    assert due.row > start.row and due.column == start.column
    assert start.border.bottom.style == "thick" and start.border.bottom.color.rgb == "00000000"
    assert due.border.top.style == "thick" and due.border.top.color.rgb == "00000000"
    assert start.border.top.color.rgb[-6:] == "DC2626"
    title = next(cell for cell in cells if str(cell.value) == "👁 Detyrë me status TODO")
    assert title.value[0].font.color.rgb[-6:] == "DC2626"
    assert title.border.top.style == title.border.bottom.style == "thick"


def test_long_answers_are_complete_and_formula_like_text_stays_literal():
    report = complete_export_fixture()
    formula_text = '=HYPERLINK("https://example.com", "ë dhe ç")'
    report["manual_answers"]["reorganization"] = formula_text
    task_id = report["data"]["delivery"][0]["task_id"]
    long_answer = "Rresht me ë dhe ç\n" * 500
    report["manual_answers"][f"delivery:{task_id}"] = long_answer
    report["data"]["delivery"][0]["title"] = "Titull\x00 i sigurt"
    sheet = load(m2.export_blocks(report))
    cells = populated(sheet)
    answer = next(cell for cell in cells if cell.value == formula_text)
    assert answer.data_type == "s"
    long_cell = next(cell for cell in cells if cell.value == long_answer)
    merged = next(area for area in sheet.merged_cells.ranges if long_cell.coordinate in area)
    assert merged.max_row > merged.min_row
    assert all(row.height <= 400 for row in sheet.row_dimensions.values())
    assert "Titull i sigurt" in [cell.value for cell in cells]
    assert not any(cell.data_type == "f" for cell in cells)


@pytest.mark.parametrize("blocks", [m2.export_blocks, m3.export_blocks])
def test_empty_reports_remain_readable(blocks):
    sheet = load(blocks({"report_date": "2026-10-07", "manual_answers": {}, "data": {}}))
    assert "Asnje detyre." in [cell.value for cell in populated(sheet)]


def test_m2_excel_export_failure_prevents_partial_email():
    import asyncio
    from app.models.m2_reporting_points import M2ReportingPointsReport
    from tests.test_m2_reporting_points import DAY

    row = M2ReportingPointsReport(report_date=DAY, manual_answers={}, data={}, status="DRAFT")
    db = AsyncMock()
    with patch.object(m2, "GmailService") as gmail, patch.object(excel, "render_xlsx", side_effect=RuntimeError("Excel export failed")):
        with pytest.raises(RuntimeError, match="Excel export failed"):
            asyncio.run(m2.send_report(db, row, {"to": ["test@example.com"]}))
    gmail.return_value.send_verified.assert_not_called()
    assert row.status == "FAILED" and row.last_error == "Excel export failed"
    db.commit.assert_awaited_once()
