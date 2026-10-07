"""Native Excel copies of the reporting view, using its existing export blocks.

The application's existing openpyxl dependency generates these on the server.
A narrow shared column grid lets each table retain its own column proportions
on one continuous worksheet, including the stacked START/DUE cells.
"""
from __future__ import annotations

import io
import math
from functools import lru_cache

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw, ImageFont

from app.services.m3_reporting_points_attachments import _text, _wrap_measured_text

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
GRID_COLUMNS = 120
GRID_PIXELS = 12


def render_xlsx(blocks: list[dict], *, sheet_name: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.freeze_panes = "A5"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.print_options.horizontalCentered = True
    for column in range(1, GRID_COLUMNS + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 1

    thin = Side(style="thin", color="000000")
    red = Side(style="thick", color="DC2626")
    divider = Side(style="thick", color="000000")
    measure = ImageDraw.Draw(Image.new("RGB", (1, 1)))

    @lru_cache(maxsize=16)
    def font(bold: bool, size: int):
        pixels = round(size * 4 / 3)
        for path in (r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
                     "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"):
            try:
                return ImageFont.truetype(path, pixels)
            except OSError:
                pass
        return ImageFont.load_default(size=pixels)

    @lru_cache(maxsize=4096)
    def height(text: str, columns: int, bold: bool = False, size: int = 10):
        lines = _wrap_measured_text(text, font(bold, size), columns * GRID_PIXELS - 14, measure)
        return max(24, len(lines) * (size + 4) + 10)

    def row_space(start: int, points: float) -> int:
        # Excel limits a physical row to 409 points. Taller content spans rows.
        count = max(1, math.ceil(points / 400))
        for row in range(start, start + count):
            sheet.row_dimensions[row].height = points / count
        return start + count - 1

    def write(start: int, end: int, left: int, right: int, value, *, fill="FFFFFF",
              color="000000", bold=False, size=10, border=None, marker=None, number_format=None):
        if isinstance(value, str):
            value = _text(value)
            if len(value) > 32767:
                raise ValueError("Teksti i raportit tejkalon kufirin e një qelize Excel.")
        cell = sheet.cell(start, left)
        cell.value = value
        # User-entered answers/titles are literal text, even if they begin '='.
        if isinstance(value, str):
            cell.data_type = "s"
        if marker and isinstance(value, str) and value.startswith(marker + " "):
            cell.value = CellRichText(
                TextBlock(InlineFont(rFont="Arial", sz=size, b=True,
                                     color="FFFFFF" if color.lstrip("#").upper() == "FFFFFF" else "DC2626"), marker + " "),
                value[len(marker) + 1:],
            )
        cell.font = Font(name="Arial", size=size, bold=bold, color=color.lstrip("#"))
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        if number_format:
            cell.number_format = number_format
        cell.border = border or Border()
        if left != right or start != end:
            sheet.merge_cells(start_row=start, end_row=end, start_column=left, end_column=right)
        # Fill the whole merged area, so all Excel viewers preserve row colors.
        paint = PatternFill("solid", fgColor=fill.lstrip("#"))
        for row in sheet.iter_rows(min_row=start, max_row=end, min_col=left, max_col=right):
            for part in row:
                part.fill = paint

    def spans(columns, compact=False):
        total = sum(width for _, _, width in columns)
        minimums = [max(3, math.ceil((measure.textlength(label, font=font(True, 10)) + 14) / GRID_PIXELS))
                    for _, label, _ in columns]
        budget = min(GRID_COLUMNS, max(sum(minimums), round(total / GRID_PIXELS))) if compact else GRID_COLUMNS
        boundaries = [0]
        used = 0
        for index, (_, _, width) in enumerate(columns[:-1]):
            used += width
            boundaries.append(max(boundaries[-1] + minimums[index],
                                  min(budget - sum(minimums[index + 1:]), round(used * budget / total))))
        boundaries.append(budget)
        return [(boundaries[i] + 1, boundaries[i + 1]) for i in range(len(columns))]

    current = 1
    for block in blocks:
        if block["kind"] == "text":
            text = _text(block["text"])
            level = block.get("level", 0)
            size = block.get("font_size", 16 if level == 1 else 11 if level else 10)
            end = row_space(current, height(text, GRID_COLUMNS, bool(level), size))
            write(current, end, 1, GRID_COLUMNS, text, bold=bool(level), size=size)
            sheet.row_dimensions[end + 1].height = 8
            current = end + 2
        elif block["kind"] == "legend":
            # Keep the same colored OK/RREZIK legend as the report view.
            for value, fill, color in block["runs"]:
                text = _text(value.strip())
                end = row_space(current, height(text, GRID_COLUMNS, True))
                write(current, end, 1, GRID_COLUMNS, text, fill=fill or "FFFFFF", color=color, bold=True)
                current = end + 1
            sheet.row_dimensions[current].height = 8
            current += 1
        else:
            columns = block["columns"]
            positions = spans(columns, block.get("compact", False))
            header_height = max(height(label, right - left + 1, True) for (_, label, _), (left, right) in zip(columns, positions))
            end = row_space(current, header_height)
            for (_, label, _), (left, right) in zip(columns, positions):
                write(current, end, left, right, label, fill="E2E8F0", bold=True,
                      border=Border(top=thin, bottom=thin, left=thin, right=thin))
            current = end + 1
            for values in block["rows"]:
                split = any(value.get("divider") for value in values)
                heights = [24, 24] if split else [24]
                for value, (left, right) in zip(values, positions):
                    text = _text(value["text"])
                    if split and value.get("divider"):
                        for index, line in enumerate(text.split("\n", 1)):
                            heights[index] = max(heights[index], height(line, right - left + 1, value.get("bold", False)))
                    else:
                        points = height(text, right - left + 1, value.get("bold", False))
                        heights = [max(points / len(heights), h) for h in heights]
                first_end = row_space(current, heights[0])
                end = row_space(first_end + 1, heights[1]) if split else first_end
                for (key, _, _), value, (left, right) in zip(columns, values, positions):
                    eight_am = value.get("eight_am")
                    top = bottom = red if eight_am else thin
                    left_edge = red if eight_am and value.get("first") else thin
                    right_edge = red if eight_am and value.get("last") else thin
                    text = _text(value["text"])
                    parts = text.split("\n", 1) if value.get("divider") and split else [text]
                    for index, part in enumerate(parts):
                        start = current if index == 0 else first_end + 1
                        stop = first_end if len(parts) == 2 and index == 0 else end
                        cell_value, number_format = part, None
                        if key == "nr":
                            cell_value = int(part)
                        elif key == "percent" and part.endswith("%"):
                            cell_value, number_format = float(part[:-1]) / 100, "0.##%"
                        write(start, stop, left, right, cell_value,
                              fill=value.get("fill", "FFFFFF"), color=value.get("color", "000000"),
                              bold=value.get("bold", False), marker=value.get("marker"), number_format=number_format,
                              border=Border(left=left_edge, right=right_edge,
                                            top=divider if index else top,
                                            bottom=divider if len(parts) == 2 and index == 0 else bottom))
                current = end + 1
            sheet.row_dimensions[current].height = 12
            current += 1

    sheet.print_area = f"A1:{get_column_letter(GRID_COLUMNS)}{current - 1}"
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
