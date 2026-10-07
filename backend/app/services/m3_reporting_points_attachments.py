"""Complete native Word and PNG exports of the M3 / GA reporting points.

Like M2/M3, these use python-docx and Pillow. They consume the same columns
and task values as the email, without converting through an ASCII table.
"""
from __future__ import annotations

import io
import os
import re
from functools import lru_cache
from typing import Any

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
GREEN, GREEN_TEXT = "#dcfce7", "#14532d"
RED, RED_TEXT = "#fee2e2", "#7f1d1d"


def _wrap_measured_text(value, font, limit, measure):
    """Keep greedy wrapping while measuring prefixes instead of every character."""
    lines = []
    for source in value.split("\n"):
        remaining = source.expandtabs(4)
        if not remaining:
            lines.append("")
            continue
        while remaining:
            if measure.textlength(remaining, font=font) <= limit:
                lines.append(remaining)
                break
            # A single glyph is retained even if it exceeds a very narrow cell.
            low, high = 1, len(remaining)
            while low < high:
                middle = (low + high + 1) // 2
                if measure.textlength(remaining[:middle], font=font) <= limit:
                    low = middle
                else:
                    high = middle - 1
            prefix = remaining[:low]
            space = prefix.rfind(" ")
            if space > 0:
                lines.append(prefix[:space])
                remaining = remaining[space + 1:]
            else:
                lines.append(prefix)
                remaining = remaining[low:]
    return tuple(lines)


def _text(value: Any) -> str:
    # XML does not accept control characters other than tab and line breaks.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))


def export_blocks(report: dict) -> list[dict]:
    from app.services.m3_reporting_points import (
        AUTO_TITLES, GA_TITLE, M3_TITLE, MANUAL_POINTS,
        report_table_columns, report_table_value, task_table_groups, task_row_appearance, realization_department_rows,
    )

    blocks: list[dict] = []

    def text(value: Any, level: int = 0):
        blocks.append({"kind": "text", "text": _text(value), "level": level})

    def manual(key: str):
        text(MANUAL_POINTS[key], 2)
        text(report.get("manual_answers", {}).get(key) or "Pa pergjigje")

    def tasks(key: str):
        text(AUTO_TITLES[key], 2)
        if key == "postponed":
            blocks.append({"kind": "legend", "runs": [
                ("Shtyrjet e bëra gjatë ditës. ", None, "#000000"),
                ("Brenda javës: OK.", GREEN, GREEN_TEXT),
                (" Për të premten ose javën tjetër: RREZIK.", RED, RED_TEXT),
            ]})
        for caption, movement, rows in task_table_groups(report.get("data", {}).get(key, []), key):
            if caption:
                text(caption + ":", 3)
            if not rows:
                text("Asnje detyre.")
                continue
            columns = report_table_columns(key, movement)
            rendered_rows = []
            for index, row in enumerate(rows, 1):
                cells = []
                row_fill, row_color, eight_am = task_row_appearance(row)
                for column_index, (column, _, _) in enumerate(columns):
                    fill = row_fill
                    color, bold = row_color, column == "title"
                    if column == "risk":
                        fill, color = (RED, RED_TEXT) if row.get("risk") == "RREZIK" else (GREEN, GREEN_TEXT)
                        bold = True
                    cells.append({"text": _text(report_table_value(row, column, index, movement)),
                                  "fill": fill, "color": color, "bold": bold,
                                  "marker": row.get("marker") if column == "title" else None,
                                  "eight_am": eight_am, "first": column_index == 0, "last": column_index == len(columns) - 1,
                                  "divider": column in {"from", "to"} and movement == "start_due"})
                rendered_rows.append(cells)
            blocks.append({"kind": "table", "columns": columns, "rows": rendered_rows})

    text(M3_TITLE, 1)
    text(f"Data: {report['report_date']}")
    for key in ("underload", "overload", "reorganization"):
        manual(key)
    for key in ("postponed", "untouched", "same_day"):
        tasks(key)
    text(AUTO_TITLES["realization"], 2)
    realization = report.get("realization")
    if realization:
        percent = realization.get("percent")
        text((f"{percent:g}%" if percent is not None else "Pa te dhena") + " — " + str(realization.get("comment") or ""))
        text(f"Marrë në: {report.get('realization_captured_at') or '16:15'}")
        department_rows = realization_department_rows(realization)
        if department_rows:
            departments = []
            for item in department_rows:
                value = item.get("percent")
                fill = "#e2e8f0" if value is None else RED if value < 50 else GREEN
                departments.append([
                    {"text": _text(item["code"]), "fill": "#ffffff", "color": "#000000", "bold": True, "divider": False},
                    {"text": f"{value:g}%" if value is not None else "Pa të dhëna", "fill": fill, "color": "#000000", "bold": True, "divider": False},
                    {"text": _text(item["comment"]), "fill": "#ffffff", "color": "#000000", "bold": False, "divider": False},
                ])
            blocks.append({"kind": "table", "columns": [("code", "DEPARTAMENTI", 124), ("percent", "REALIZIMI", 96), ("comment", "VLERËSIMI", 360)], "rows": departments})
    else:
        text("Vlera e realizimit merret ne 16:15. Nuk ka vlere te ruajtur per kete date.")
    text(GA_TITLE, 1)
    manual("ga_reorganization")
    tasks("ga_postponed")
    return blocks


def render_docx(report: dict, *, blocks: list[dict] | None = None) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor, Twips

    document = Document()
    page = document.sections[0]
    # A3 landscape keeps all utility, date, title and comment columns legible.
    page.orientation = WD_ORIENT.LANDSCAPE
    page.page_width, page.page_height = Inches(16.54), Inches(11.69)
    page.top_margin = page.bottom_margin = Inches(0.5)
    page.left_margin = page.right_margin = Inches(0.6)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(9)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(4)
    if blocks is not None:
        # Custom report titles use plain black typography, without the default
        # Word template's blue paragraph border.
        for border in list(document.styles["Title"]._element.xpath("./w:pPr/w:pBdr")):
            border.getparent().remove(border)
    available = int((page.page_width - page.left_margin - page.right_margin) / 635)

    def font(run, *, bold=False, color="#000000", size=8):
        run.font.name, run.font.size = "Arial", Pt(size)
        run.bold = bold
        run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))

    def shade(properties, fill):
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill.lstrip("#"))
        properties.append(shading)

    def write_cell(cell, value, fill, color="#000000", bold=False, divider=False, eight_am=False, first=False, last=False, marker=None):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        properties = cell._tc.get_or_add_tcPr()
        shade(properties, fill)
        borders = OxmlElement("w:tcBorders")
        for name in ("top", "left", "bottom", "right"):
            border = OxmlElement(f"w:{name}")
            red_edge = eight_am and (name in {"top", "bottom"} or name == "left" and first or name == "right" and last)
            for attr, item in (("val", "single"), ("sz", "18" if red_edge else "6"), ("color", "DC2626" if red_edge else "000000")):
                border.set(qn(f"w:{attr}"), item)
            borders.append(border)
        properties.append(borders)
        lines = value.split("\n", 1) if divider else [value]
        for index, line in enumerate(lines):
            paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(2 if divider else 0)
            paragraph.paragraph_format.keep_with_next = divider and index == 0
            if marker and line.startswith(marker + " "):
                marker_run = paragraph.add_run(marker + " ")
                font(marker_run, bold=True, color="#ffffff" if color == "#ffffff" else "#dc2626", size=10)
                if "👁" in marker:
                    marker_run.font.name = "Segoe UI Symbol"
                font(paragraph.add_run(line[len(marker) + 1:]), bold=bold, color=color)
            else:
                font(paragraph.add_run(line), bold=bold, color=color)
            if divider and index == 0:
                paragraph_borders = OxmlElement("w:pBdr")
                bottom = OxmlElement("w:bottom")
                for attr, item in (("val", "single"), ("sz", "18"), ("space", "2"), ("color", "000000")):
                    bottom.set(qn(f"w:{attr}"), item)
                paragraph_borders.append(bottom)
                paragraph._p.get_or_add_pPr().append(paragraph_borders)

    for block in (export_blocks(report) if blocks is None else blocks):
        if block["kind"] == "text":
            level = block["level"]
            paragraph = document.add_paragraph(style="Title" if level == 1 else None)
            paragraph.paragraph_format.space_before = Pt(8 if level else 0)
            paragraph.paragraph_format.keep_with_next = bool(level)
            font(paragraph.add_run(block["text"]), bold=bool(level), size=block.get("font_size", {1: 16, 2: 10, 3: 9}.get(level, 9)))
        elif block["kind"] == "legend":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.keep_with_next = True
            for value, fill, color in block["runs"]:
                run = paragraph.add_run(value)
                font(run, bold=True, color=color, size=9)
                if fill:
                    shade(run._element.get_or_add_rPr(), fill)
        else:
            columns = block["columns"]
            total = sum(column[2] for column in columns)
            widths = [int(available * column[2] / total) for column in columns]
            widths[-1] += available - sum(widths)
            table = document.add_table(rows=1, cols=len(columns))
            table.autofit = False
            for index, width in enumerate(widths):
                table.columns[index].width = Twips(width)
            repeat = OxmlElement("w:tblHeader")
            table.rows[0]._tr.get_or_add_trPr().append(repeat)
            for index, (_, label, _) in enumerate(columns):
                cell = table.rows[0].cells[index]
                cell.width = Twips(widths[index])
                write_cell(cell, label, "#e2e8f0", bold=True)
            for values in block["rows"]:
                row = table.add_row()
                if blocks is not None:
                    # Keep an ordinary M2 task row together across page breaks.
                    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
                for index, value in enumerate(values):
                    cell = row.cells[index]
                    cell.width = Twips(widths[index])
                    write_cell(cell, **{key: value[key] for key in ("fill", "color", "bold", "divider", "eight_am", "first", "last", "marker") if key in value}, value=value["text"])
            document.add_paragraph().paragraph_format.space_after = Pt(3)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def render_png(report: dict, *, blocks: list[dict] | None = None) -> bytes:
    from PIL import Image, ImageDraw, ImageFont

    blocks = export_blocks(report) if blocks is None else blocks
    # At least the email's full column budget, enlarged for readable text.
    scale, margin = 1.4, 40
    content_width = max(1400, round(max((sum(c[2] for c in b["columns"]) for b in blocks if b["kind"] == "table"), default=1000) * scale))
    width = content_width + 2 * margin

    def load_font(bold, size):
        candidates = [os.getenv("PRIMEFLOW_REPORT_FONT_PATH", "") if not bold else "",
                      r"C:\Windows\Fonts\arialbd.ttf" if bold else r"C:\Windows\Fonts\arial.ttf",
                      "DejaVuSans-Bold.ttf" if bold else "DejaVuSans.ttf"]
        for candidate in candidates:
            if candidate:
                try:
                    return ImageFont.truetype(candidate, size)
                except OSError:
                    pass
        return ImageFont.load_default(size=size)

    fonts = {"normal": load_font(False, 17), "bold": load_font(True, 17), "title": load_font(True, 28)}
    measure = ImageDraw.Draw(Image.new("RGB", (1, 1)))

    @lru_cache(maxsize=4096)
    def wrap(value, font, limit):
        # Repeated initials, dates and comments share measurements in this PNG.
        # The cache is discarded with the render; report content is not retained.
        return _wrap_measured_text(value, font, limit, measure)

    layout = []
    for block in blocks:
        if block["kind"] == "text":
            text_font = fonts["title" if block["level"] == 1 else "bold" if block["level"] else "normal"]
            step = 36 if block["level"] == 1 else 25
            if block.get("font_size"):
                text_font = load_font(True, round(block["font_size"] * 1.7))
                step = round(block["font_size"] * 2.2)
            lines = wrap(block["text"], text_font, content_width)
            layout.append({"kind": "text", "lines": lines, "font": text_font, "step": step, "height": len(lines) * step + 14})
        elif block["kind"] == "legend":
            # Separate lines keep the meaning and colored highlights legible.
            for value, fill, color in block["runs"]:
                lines = wrap(value.strip(), fonts["bold"], content_width - 12)
                layout.append({"kind": "text", "lines": lines, "font": fonts["bold"], "step": 25,
                               "height": len(lines) * 25 + 6, "fill": fill, "color": color})
        else:
            total = sum(c[2] for c in block["columns"])
            widths = [round(content_width * c[2] / total) for c in block["columns"]]
            widths[-1] += content_width - sum(widths)
            header = [{"text": label, "fill": "#e2e8f0", "color": "#000000", "bold": True, "divider": False}
                      for _, label, _ in block["columns"]]
            for values in [header, *block["rows"]]:
                cells = []
                for index, value in enumerate(values):
                    font = fonts["bold" if value["bold"] else "normal"]
                    lines = wrap(value["text"], font, max(10, widths[index] - 12))
                    divider_after = len(wrap(value["text"].split("\n", 1)[0], font, max(10, widths[index] - 12))) if value["divider"] else None
                    cells.append({**value, "lines": lines, "font": font, "divider_after": divider_after})
                height = max(len(cell["lines"]) * 24 + 14 + (8 if cell["divider"] else 0) for cell in cells)
                layout.append({"kind": "row", "cells": cells, "widths": widths, "height": height})
            layout.append({"kind": "space", "height": 18})
    height = 2 * margin + sum(item["height"] for item in layout)
    image = Image.new("RGB", (width, height), "white")
    draw, y = ImageDraw.Draw(image), margin
    for block in layout:
        if block["kind"] == "text":
            for index, line in enumerate(block["lines"]):
                line_y = y + index * block["step"]
                if block.get("fill"):
                    draw.rectangle((margin, line_y, margin + draw.textlength(line, font=block["font"]) + 10, line_y + block["step"]), fill=block["fill"])
                draw.text((margin + (4 if block.get("fill") else 0), line_y), line, font=block["font"], fill=block.get("color", "#000000"))
        elif block["kind"] == "row":
            x = margin
            for cell, cell_width in zip(block["cells"], block["widths"]):
                right = x + cell_width
                draw.rectangle((x, y, right, y + block["height"]), fill=cell["fill"], outline="black")
                for index, line in enumerate(cell["lines"]):
                    after_divider = cell["divider_after"] is not None and index >= cell["divider_after"]
                    text_y = y + 6 + index * 24 + (8 if after_divider else 0)
                    marker = cell.get("marker")
                    if index == 0 and marker and line.startswith(marker + " "):
                        marker_text = marker + " "
                        marker_color = "#ffffff" if cell["color"] == "#ffffff" else "#dc2626"
                        if "👁" in marker:
                            # Draw the eye directly: Pillow has no font fallback
                            # for this emoji on every deployed Windows server.
                            prefix, suffix = marker.split("👁", 1)
                            draw.text((x + 6, text_y), prefix, font=cell["font"], fill=marker_color)
                            eye_x = x + 6 + draw.textlength(prefix, font=cell["font"])
                            eye_width = draw.textlength("👁", font=cell["font"])
                            draw.ellipse((eye_x, text_y + 7, eye_x + eye_width, text_y + 16), outline=marker_color, width=2)
                            draw.ellipse((eye_x + eye_width / 2 - 2, text_y + 9, eye_x + eye_width / 2 + 2, text_y + 14), fill=marker_color)
                            draw.text((eye_x + eye_width, text_y), suffix, font=cell["font"], fill=marker_color)
                        else:
                            draw.text((x + 6, text_y), marker_text, font=cell["font"], fill=marker_color)
                        draw.text((x + 6 + draw.textlength(marker_text, font=cell["font"]), text_y), line[len(marker_text):], font=cell["font"], fill=cell["color"])
                    else:
                        draw.text((x + 6, text_y), line, font=cell["font"], fill=cell["color"])
                if cell["divider_after"] is not None:
                    line_y = y + 6 + cell["divider_after"] * 24
                    draw.line((x + 3, line_y, right - 3, line_y), fill="black", width=3)
                x = right
            if any(cell.get("eight_am") for cell in block["cells"]):
                draw.rectangle((margin, y, margin + content_width, y + block["height"]), outline="#dc2626", width=3)
        y += block["height"]
    output = io.BytesIO()
    # Default lossless compression avoids the extra optimization pass. Pixels
    # remain identical; files grow slightly while encoding finishes much sooner.
    image.save(output, format="PNG", compress_level=6)
    return output.getvalue()


def report_attachments(report: dict) -> list[tuple[str, bytes, str]]:
    filename = f"PrimeFlow-PIKAT-M3-GA-{report['report_date']}"
    return [(filename + ".docx", render_docx(report), DOCX_MIME),
            (filename + ".png", render_png(report), "image/png")]
